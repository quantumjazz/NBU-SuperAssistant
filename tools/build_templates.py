"""Build the three Word templates in templates/ from the owner's samples.

Development only; the service never runs this. Usage, from the repository:

    python tools/build_templates.py /path/to/examples

The samples hold real students' data and stay outside the repository
(CLAUDE.md "Data in the repository"). This script reads them in memory, keeps
each one's layout (page, margins, fonts, table borders) by changing only the
text of the runs it fills, and writes:

- templates/zaiavlenie.docx   from the blank заявление-декларация form;
- templates/prilozhenie2.docx from the blank Приложение 2 form;
- templates/doklad_temi.docx  from the bachelor's topic-approval report, with
  its header row, one templated row between docxtpl's two {%tr %} control
  rows, and every student row removed.

Each dotted placeholder the records fill becomes one docxtpl tag; the text
behind each tag is worked out in superassistant/documents.py, so the
templates hold no logic (CLAUDE.md "Documents"). The years the samples print
are removed, so a date left blank is blank in full (NOW.md requirement 4),
and the core properties (author, last modified by, comments, dates) are
cleared to neutral values.

Last, it proves that no student survived: every text of the report samples'
student rows is compared with every text node of every part of the three
templates, and only the counts are printed, never a student's text.
"""

import copy
import io
import re
import sys
import zipfile
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.text.paragraph import Paragraph
from docx.text.run import Run

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / "templates"

ZAIAVLENIE_SAMPLE = "zaiavlenie-za-diplomirane-11-09-2024.docx"
PRILOZHENIE_SAMPLE = "prilojenie-2-2025-n-doc.docx"
REPORT_SAMPLES = (
    "Доклад утвърждаване на теми БФ - ПА - септември 2026.docx",
    "Доклад утвърждаване на теми МФ - секция ПА – Септември.docx",
)

NEUTRAL_CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/'
    '2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/'
    '1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://'
    'purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-'
    'instance"><dc:title></dc:title><dc:subject></dc:subject><dc:creator>'
    '</dc:creator><cp:keywords></cp:keywords><dc:description>'
    '</dc:description><cp:lastModifiedBy></cp:lastModifiedBy><cp:revision>1'
    '</cp:revision></cp:coreProperties>')

NEUTRAL_CUSTOM = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/'
    'custom-properties" xmlns:vt="http://schemas.openxmlformats.org/'
    'officeDocument/2006/docPropsVTypes"></Properties>')

EMPTY_RELS = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
    'relationships"></Relationships>')


class BuildError(Exception):
    """A sample no longer looks the way this script expects."""


# --- Editing paragraphs and runs ---------------------------------------------

def body_paragraphs(document):
    return list(document.paragraphs)


def find(paragraphs, prefix, exact=False):
    """The one paragraph whose text starts with prefix (or is exactly it)."""
    found = [p for p in paragraphs
             if (p.text == prefix if exact else p.text.startswith(prefix))]
    if len(found) != 1:
        raise BuildError("expected one paragraph starting {!r}, found {}"
                         .format(prefix, len(found)))
    return found[0]


def set_runs(paragraph, texts):
    """Give the paragraph's runs new texts: texts maps a run's index to its
    new text; every run not named is removed. Each kept run keeps its own
    formatting, so the sample's fonts survive."""
    runs = paragraph.runs
    for index, run in enumerate(runs):
        if index in texts:
            run.text = texts[index]
        else:
            run._r.getparent().remove(run._r)


def drop_bookmarks(element):
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for node in element.findall(".//" + qn(tag)):
            node.getparent().remove(node)


def new_paragraph_after(paragraph, model, text):
    """A copy of model's paragraph, holding only text, put after paragraph."""
    element = copy.deepcopy(model._p)
    runs = element.findall(qn("w:r"))
    for run in runs[1:]:
        element.remove(run)
    paragraph._p.addnext(element)
    new = Paragraph(element, paragraph._parent)
    new.runs[0].text = text
    return new


# --- The three templates -----------------------------------------------------

def build_zaiavlenie(path):
    document = docx.Document(str(path))
    paragraphs = body_paragraphs(document)
    set_runs(find(paragraphs, "Трите имена:"),
             {0: "Трите имена: {{ names }}"})
    set_runs(find(paragraphs, "F…"),
             {0: "{{ fn }}", 1: " тел.: {{ phone }} ", 2: "e-mail:",
              3: " {{ email }}"})
    set_runs(find(paragraphs, "Програма:"), {0: "Програма: {{ programme }}"})
    set_runs(find(paragraphs, "Моля да бъда допуснат(а)"),
             {0: "Моля да бъда допуснат(а) до дипломиране чрез "
                 "{{ grad_type }}"})

    exam_table, thesis_table = document.tables
    exam = exam_table.cell(0, 0).paragraphs
    set_runs(find(exam, "Дата на подаване"),
             {0: "Дата на подаване на заявлението: {{ exam_date }}"
                 "         Подпис: …………………………………."})

    thesis = thesis_table.cell(0, 0).paragraphs
    lines = [p for p in thesis if p.text.startswith("……")]
    if len(lines) != 3:
        raise BuildError("expected three dotted topic lines")
    for paragraph, tag in zip(lines, ("topic", "topic_line2",
                                      "topic_line3")):
        set_runs(paragraph, {0: "{{ " + tag + " }}"})
    set_runs(find(thesis, "С научен (художествен) ръководител:"),
             {0: "С научен (художествен) ръководител: {{ supervisor }}"})
    set_runs(find(thesis, "среден успех от следването:"),
             {0: "среден успех от следването: {{ average_grade }}"})
    set_runs(find(thesis, "Дата на подаване"),
             {0: "Дата на подаване на заявлението: {{ thesis_date }}"
                 "         ", 1: "Подпис: …………………………………."})
    return document


PROPOSAL_HEADINGS = (
    ("ОБОСНОВКА НА ТЕМАТА", "rationale"),
    ("ЦЕЛИ И ЗАДАЧИ", "aims"),
    ("ОБЕКТ И ПРЕДМЕТ НА ИЗСЛЕДВАНЕТО", "object_subject"),
    ("МЕТОДОЛОГИЧНА БАЗА НА ИЗСЛЕДВАНЕТО", "methodology"),
    ("ИНФОРМАЦИОННА БАЗА НА ИЗСЛЕДВАНЕТО", "information_base"),
)


def is_dots(text):
    return bool(text.strip()) and not text.strip(" .…")


def build_prilozhenie(path):
    document = docx.Document(str(path))
    paragraphs = body_paragraphs(document)

    # The sample's title mixes two Latin letters into "ПРЕДЛОЖЕНИЕ"
    # (docs/domain.md, "Errors found in the samples"): written in Cyrillic.
    title = find(paragraphs, "ПPEДЛОЖЕНИЕ")
    for run in title.runs:
        if run.text == "PE":
            run.text = "РЕ"
    heading = find(paragraphs, "БАКАЛАВЪРСКА/ МАГИСТЪРСКА")
    set_runs(heading, {0: "{{ thesis_heading }}", 1: " ", 2: "ТЕЗА"})

    student = find(paragraphs, "На студента")
    if [r.text for r in student.runs][7] != "F":
        raise BuildError("the student's line has changed")
    set_runs(student, {0: "На", 1: " ", 2: "студента", 3: " ",
                       4: "{{ names }}", 6: " ", 7: "{{ fn }}"})

    topic = find(paragraphs, "ТЕМА")
    set_runs(topic, {0: "ТЕМА", 1: " {{ topic }}"})

    plan_heading = find(paragraphs, "ПРИМЕРЕН ПЛАН-СЪДЪРЖАНИЕ")
    plan_line = paragraphs[paragraphs.index(plan_heading) + 1]
    schedule_heading = find(paragraphs, "ВРЕМЕВИ ГРАФИК")
    schedule_line = paragraphs[paragraphs.index(schedule_heading) + 1]
    if not (is_dots(plan_line.text) and is_dots(schedule_line.text)):
        raise BuildError("the plan and schedule lines have changed")
    # The text of each section goes on its own paragraph after its heading,
    # formatted like the sample's own dotted lines for the plan.
    model = copy.deepcopy(plan_line._p)
    model_paragraph = Paragraph(model, plan_line._parent)
    for prefix, tag in PROPOSAL_HEADINGS:
        paragraph = find(paragraphs, prefix)
        keep = {}
        runs = paragraph.runs
        # Drop the dotted runs at the end, and the space before them.
        last = len(runs)
        while last and (is_dots(runs[last - 1].text)
                        or not runs[last - 1].text.strip()):
            last -= 1
        for index in range(last):
            keep[index] = runs[index].text
        set_runs(paragraph, keep)
        new_paragraph_after(paragraph, model_paragraph, "{{ " + tag + " }}")
    set_runs(plan_line, {0: "{{ plan }}"})
    set_runs(schedule_line, {0: "{{ schedule }}"})
    # The one heading in the sample that is not in the "Heading 1" style got
    # its space from the heading above it; with a section's text now between
    # them, it gets that style's 5 pt before it, like the other headings.
    find(paragraphs, "МЕТОДОЛОГИЧНА БАЗА").paragraph_format.space_before = \
        Pt(5)

    handed_in = find(paragraphs, "Дата на предаване:")
    texts = {i: r.text for i, r in enumerate(handed_in.runs)}
    year = [i for i, text in texts.items() if re.search(r"20\d\d", text)]
    if len(year) != 1 or texts[1] != "…………….":
        raise BuildError("the date line has changed")
    texts[1] = "………………………"
    del texts[year[0]]
    set_runs(handed_in, texts)
    return document


REPORT_COLUMNS = ("{{ s.position }}", "{{ s.fn }}", "{{ s.name }}",
                  "{{ s.topic }}", "{{ s.supervisor }}", None)
HEADER = ["№", "F№", "Име", "Тема", "Научен ръководител", "Рецензент"]


def emptied_row(row_element):
    """A copy of a table row with every run and bookmark removed."""
    row = copy.deepcopy(row_element)
    drop_bookmarks(row)
    for cell in row.findall(qn("w:tc")):
        paragraphs = cell.findall(qn("w:p"))
        for extra in paragraphs[1:]:
            cell.remove(extra)
        for run in paragraphs[0].findall(qn("w:r")):
            paragraphs[0].remove(run)
    return row


def put_text(cell_element, model_run, text):
    paragraph = cell_element.find(qn("w:p"))
    run = copy.deepcopy(model_run)
    for child in list(run):
        if child.tag != qn("w:rPr"):
            run.remove(child)
    paragraph.append(run)
    Run(run, None).text = text


def build_report(path):
    document = docx.Document(str(path))
    paragraphs = body_paragraphs(document)

    dean = paragraphs[paragraphs.index(find(paragraphs, "ДО", exact=True)) + 1]
    set_runs(dean, {0: "{{ dean }}"})
    set_runs(find(paragraphs, "ДЕКАН НА"), {0: "ДЕКАН НА",
                                            1: " {{ faculty }}"})
    number = find(paragraphs, "Изх. №")
    if [r.text for r in number.runs] != ["Изх. № …….", " /", "……..", ".",
                                         "202", "6", " г."]:
        raise BuildError("the outgoing-number line has changed")
    set_runs(number, {0: "Изх. № …….", 1: " /", 2: "……..", 3: ".",
                      4: "…", 6: " г."})

    intro = find(paragraphs, "Студент от")
    runs = intro.runs
    if not (runs[1].bold and runs[3].bold and not runs[0].bold
            and not runs[7].bold):
        raise BuildError("the introductory sentence has changed")
    set_runs(intro, {
        0: "Студенти от ",
        1: "{{ degree }} „{{ programme }}“",
        7: ", изпълнили условията за допускане до дипломиране с "
           "{{ thesis_adjective }} теза:"})

    signature = find(paragraphs, " \t")
    for run in signature.runs:
        if run.text == "доц. д-р  ":
            run.text = "доц. д-р "

    dated = find(paragraphs, "Дата:")
    texts = {i: r.text for i, r in enumerate(dated.runs)}
    if texts[1] != "ата:  " or not re.search(r"20\d\d", texts[3]):
        raise BuildError("the date line has changed")
    texts[1] = "ата: "
    texts[2] = "{{ date }}"
    del texts[3]
    set_runs(dated, texts)

    (table,) = document.tables
    if [c.text for c in table.rows[0].cells] != HEADER:
        raise BuildError("the report's columns have changed")
    header_to_12pt(table)
    rows = table._tbl.findall(qn("w:tr"))
    model = rows[1]
    model_run = model.find(".//" + qn("w:r"))
    data = emptied_row(model)
    for cell, tag in zip(data.findall(qn("w:tc")), REPORT_COLUMNS):
        if tag:
            put_text(cell, model_run, tag)
    opening = emptied_row(model)
    closing = emptied_row(model)
    put_text(opening.find(qn("w:tc")), model_run,
             "{%tr for s in students %}")
    put_text(closing.find(qn("w:tc")), model_run, "{%tr endfor %}")
    for row in rows[1:]:
        table._tbl.remove(row)
    for row in (opening, data, closing):
        table._tbl.append(row)
    return document


def header_to_12pt(table):
    """The header row in the adopted layout's size: Times New Roman 12 pt
    (docs/domain.md). The samples set it at 10 pt."""
    for cell in table.rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(12)
            marks = paragraph._p.find(qn("w:pPr"))
            if marks is not None and marks.find(qn("w:rPr")) is not None:
                for name in ("w:sz", "w:szCs"):
                    size = marks.find(qn("w:rPr")).find(qn(name))
                    if size is not None:
                        size.set(qn("w:val"), "24")


# --- Saving, with neutral metadata ---------------------------------------------

def save(document, target):
    buffer = io.BytesIO()
    document.save(buffer)
    source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as package:
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                data = NEUTRAL_CORE.encode("utf-8")
            elif info.filename == "docProps/custom.xml":
                data = NEUTRAL_CUSTOM.encode("utf-8")
            elif info.filename == "word/_rels/settings.xml.rels":
                # An attached template on someone's own disk.
                data = EMPTY_RELS.encode("utf-8")
            elif info.filename == "word/settings.xml":
                data = re.sub(rb"<w:attachedTemplate [^>]*/>", b"", data)
            package.writestr(info.filename, data)
    target.write_bytes(out.getvalue())


# --- The proof that no student survived --------------------------------------

def texts_of_parts(path):
    """Every text node of every XML part of a .docx, and each part's text
    joined, so that a text split across runs is found too."""
    nodes, joined = [], []
    with zipfile.ZipFile(str(path)) as package:
        for name in package.namelist():
            if not (name.endswith(".xml") or name.endswith(".rels")):
                continue
            xml = package.read(name).decode("utf-8")
            found = [t for t in re.findall(r">([^<]+)<", xml)]
            nodes.extend(found)
            joined.append("".join(found))
            joined.append(" ".join(found))
    return nodes, joined


def student_texts(examples):
    """The texts of the report samples' student rows, in memory only: each
    cell and each paragraph in it, but not the № column, and not a text the
    sample also prints outside its student rows (a supervisor who is also
    the report's author is not a student's data)."""
    student, elsewhere = set(), set()
    rows = 0
    for filename in REPORT_SAMPLES:
        document = docx.Document(str(examples / filename))
        for paragraph in document.paragraphs:
            elsewhere.add(paragraph.text.strip())
        for table in document.tables:
            for cell in table.rows[0].cells:
                elsewhere.add(cell.text.strip())
            for row in table.rows[1:]:
                rows += 1
                for cell in row.cells[1:]:
                    student.add(cell.text.strip())
                    for paragraph in cell.paragraphs:
                        student.add(paragraph.text.strip())
    for text in list(elsewhere):
        for part in re.split(r"\s{2,}|\t", text):
            elsewhere.add(part.strip())
    student = {t for t in student if t and t not in elsewhere}
    return student, rows


def prove_no_student(examples, built):
    student, rows = student_texts(examples)
    nodes, joined = [], []
    for path in built:
        n, j = texts_of_parts(path)
        nodes.extend(n)
        joined.extend(j)
    node_set = set(t.strip() for t in nodes)
    exact = sum(1 for t in student if t in node_set)
    inside = sum(1 for t in student if len(t) >= 4
                 and any(t in part for part in joined))
    print("Student rows read from the samples: {}".format(rows))
    print("Distinct student texts compared: {}".format(len(student)))
    print("Template text nodes compared: {}".format(len(nodes)))
    print("Student texts equal to a text node: {}".format(exact))
    print("Student texts found inside a template's text: {}".format(inside))
    return exact == 0 and inside == 0


def prove_no_year(built):
    found = 0
    for path in built:
        nodes, _ = texts_of_parts(path)
        with zipfile.ZipFile(str(path)) as package:
            for name in package.namelist():
                if name.startswith("word/") and name.endswith(".xml"):
                    xml = package.read(name).decode("utf-8")
                    for text in re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>",
                                           xml):
                        if re.search(r"20\d\d", text):
                            found += 1
    print("Text nodes printing a year 20xx: {}".format(found))
    return found == 0


def main(argv):
    if len(argv) != 2:
        sys.stderr.write("usage: python tools/build_templates.py "
                         "<examples folder>\n")
        return 2
    examples = Path(argv[1])
    TEMPLATES.mkdir(exist_ok=True)
    built = []
    for builder, sample, name in (
            (build_zaiavlenie, ZAIAVLENIE_SAMPLE, "zaiavlenie.docx"),
            (build_prilozhenie, PRILOZHENIE_SAMPLE, "prilozhenie2.docx"),
            (build_report, REPORT_SAMPLES[0], "doklad_temi.docx")):
        target = TEMPLATES / name
        save(builder(examples / sample), target)
        built.append(target)
        print("Wrote templates/" + name)
    clean = prove_no_student(examples, built)
    clean = prove_no_year(built) and clean
    if not clean:
        sys.stderr.write("A template still holds text it must not hold.\n")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
