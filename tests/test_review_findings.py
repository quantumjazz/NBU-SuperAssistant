"""Regressions for Codex's review of milestone 1 at aa278fc (M1-C3 to M1-C9;
M1-C1 and M1-C2 are browser regressions in tests/browser).

Invented people, temporary data and the fake relay only.
"""

import io
import re

import pytest
from docx import Document
from hypothesis import HealthCheck, given, settings as hsettings
from hypothesis import strategies as st

from superassistant import applications, db, limits, lists, reports
from superassistant.web import StartupError, create_app
from tests.conftest import (Client, connect, make_account, make_settings,
                            prepare_data_dir)
from tests.test_application_flow import (THESIS_TAB_1, THESIS_TAB_2,
                                         as_staff, as_student, link_from,
                                         record_for, save, submit_thesis)


# --- M1-C4: staff correct the graduation type -------------------------------

def submit_exam(client, address, extra=None):
    as_student(client, address)
    fields = {"name1": "Калина", "name2": "Иванова", "name3": "Стоева",
              "fn": "500001", "grad_type": "exam", "credits_declared": "1"}
    fields.update(extra or {})
    assert save(client, 1, fields, "submit").status_code == 302


def thesis_data():
    data = dict(THESIS_TAB_1)
    data.update(THESIS_TAB_2)
    data.pop("grad_type")
    data.pop("professor")
    data.pop("credits_declared")
    return data


@pytest.fixture
def staff(settings):
    return make_account(settings, "staff1")


def test_exam_to_thesis_needs_the_thesis_data_first(client, new_client,
                                                    settings, outbox, staff):
    submit_exam(client, "type1@example.org")
    record = record_for(settings, "type1@example.org")
    viewer = as_staff(new_client(), staff)
    base = "/staff/application/{}".format(record["id"])
    response = viewer.post(base + "/type", data={"grad_type": "thesis",
                                                 "professor": "ivanov"})
    assert response.status_code == 400
    assert "За теза липсват" in response.get_data(as_text=True)
    assert record_for(settings, "type1@example.org")["grad_type"] == "exam"
    connection = connect(settings)
    assert applications.corrections_of(connection, record["id"]) == []
    connection.close()
    assert outbox.messages == []


def test_exam_to_thesis_and_back_keeps_history_and_answers(
        client, new_client, settings, outbox, staff):
    submit_exam(client, "type2@example.org")
    record = record_for(settings, "type2@example.org")
    viewer = as_staff(new_client(), staff)
    base = "/staff/application/{}".format(record["id"])
    for field, value in thesis_data().items():
        assert viewer.post(base + "/correct", data={
            "field": field, "value": value}).status_code == 302
    outbox.clear()
    # The relay refuses: nothing changes.
    outbox.refuse = True
    assert viewer.post(base + "/type", data={
        "grad_type": "thesis", "professor": "ivanov"}).status_code == 503
    assert record_for(settings, "type2@example.org")["grad_type"] == "exam"
    outbox.refuse = False
    assert viewer.post(base + "/type", data={
        "grad_type": "thesis", "professor": "ivanov"}).status_code == 302
    assert record_for(settings, "type2@example.org")["grad_type"] == "thesis"
    (request_mail,) = outbox.messages
    assert request_mail.intended == "p.ivanov@example.org"
    connection = connect(settings)
    type_rows = [c for c in applications.corrections_of(connection,
                                                        record["id"])
                 if c["field"] == "grad_type"]
    assert [(c["old_value"], c["new_value"], c["username"])
            for c in type_rows] == [("exam", "thesis", "staff1")]
    connection.close()
    # The student's page and forms follow the corrected type.
    page = client.get("/application").get_data(as_text=True)
    assert "Чака отговор" in page and "prilozhenie2" in page
    assert client.get("/application/forms/prilozhenie2.docx").status_code \
        == 200
    # Eligible for the report only once confirmed and checked.
    connection = connect(settings)
    loaded = lists.load(settings.lists_dir)
    qualifying = lambda: [r["id"] for r in reports.qualifying(connection,
                                                              "BSNBUB04")]
    assert record["id"] not in qualifying()
    new_client().post(link_from(request_mail, "confirm"), data={})
    applications.record_check(connection, record["id"], "passed", "",
                              staff["id"])
    assert record["id"] in qualifying()
    # Back to a state exam: the request standing is retired, its answer kept.
    assert viewer.post(base + "/type",
                       data={"grad_type": "exam"}).status_code == 302
    rows = applications.requests_of(connection, record["id"])
    assert [r["answer"] for r in rows] == ["confirm"]
    assert rows[0]["replaced_at"] is not None
    assert record["id"] not in qualifying()
    connection.close()
    assert client.get("/application/forms/prilozhenie2.docx").status_code \
        == 404
    page = client.get("/application").get_data(as_text=True)
    assert "prilozhenie2" not in page and "Чака отговор" not in page
    staff_page = viewer.get(base).get_data(as_text=True)
    assert "държавен изпит" in staff_page and "теза" in staff_page
    assert loaded is not None


# --- M1-C5: a maintained list is never recreated -----------------------------

def test_a_list_removed_after_installation_stops_the_start(tmp_path):
    data_dir = tmp_path / "data"
    settings = make_settings(data_dir)
    create_app(settings)                          # first installation
    path = settings.lists_dir / lists.PROGRAMMES
    text = path.read_text(encoding="utf-8").replace(
        "Администрация и управление", "Публична администрация")
    path.write_text(text, encoding="utf-8")
    from tests.conftest import insert_application
    insert_application(settings, "c5@example.org", programme="BSNADM01")
    path.unlink()
    with pytest.raises(StartupError) as raised:
        create_app(settings)
    assert lists.PROGRAMMES in str(raised.value)
    assert "missing" in str(raised.value)
    assert not path.exists()                      # not quietly recreated


# --- M1-C6: list text a document cannot hold --------------------------------

def professors_text(rows):
    return "код,звание,име,email\n" + "".join(row + "\n" for row in rows)


@hsettings(max_examples=60, deadline=None,
           suppress_health_check=[HealthCheck.function_scoped_fixture])
@given(before=st.integers(min_value=0, max_value=4),
       multiline=st.booleans(),
       bad=st.characters(whitelist_categories=("Cc",)).filter(
           lambda ch: ch not in "\n\r"),
       position=st.integers(min_value=0, max_value=5))
def test_an_unprintable_character_is_reported_by_its_true_line(
        tmp_path_factory, before, multiline, bad, position):
    rows = ['p{0},доц. д-р,Име Фамилия{0},p{0}@example.org'.format(i)
            for i in range(before)]
    if multiline:
        rows.insert(0, 'm,доц. д-р,"Две\nреда",m@example.org')
    name = "Петър Иванов"
    position = min(position, len(name))
    rows.append('x,доц. д-р,{},x@example.org'.format(
        name[:position] + bad + name[position:]))
    text = professors_text(rows)
    expected_line = text.split("\n").index(next(
        line for line in text.split("\n") if bad in line)) + 1
    folder = tmp_path_factory.mktemp("lists")
    lists.write_starter(folder)
    (folder / lists.PROFESSORS).write_text(text, encoding="utf-8")
    with pytest.raises(lists.ListError) as raised:
        lists.load(folder)
    assert raised.value.line == expected_line
    assert "U+{:04X}".format(ord(bad)) in str(raised.value)


def test_a_line_break_inside_a_cell_is_refused_and_later_lines_stay_true(
        tmp_path):
    lists.write_starter(tmp_path)
    (tmp_path / lists.PROFESSORS).write_text(professors_text([
        'a,доц. д-р,"Едно\nдве",a@example.org']), encoding="utf-8")
    with pytest.raises(lists.ListError) as raised:
        lists.load(tmp_path)
    assert raised.value.line == 2 and "line break" in str(raised.value)
    (tmp_path / lists.PROFESSORS).write_text(professors_text([
        'a,доц. д-р,Едно Две,a@example.org',
        'b,акад.,Три Четири,b@example.org']), encoding="utf-8")
    with pytest.raises(lists.ListError) as raised:
        lists.load(tmp_path)
    assert raised.value.line == 3


def test_a_nul_in_a_professor_stops_the_start(tmp_path):
    data_dir = tmp_path / "data"
    prepare_data_dir(data_dir)
    (data_dir / "lists" / lists.PROFESSORS).write_text(professors_text([
        "ivanov,доц. д-р,Пе\x00тър Иванов,p.ivanov@example.org"]),
        encoding="utf-8")
    with pytest.raises(StartupError) as raised:
        create_app(make_settings(data_dir))
    assert "professors.csv line 2" in str(raised.value)
    assert "U+0000" in str(raised.value)


# --- M1-C7: a forwarded header does not reset the source limit ---------------

def test_a_spoofed_forwarding_header_does_not_bypass_the_source_limit(
        app, outbox):
    statuses = []
    for number in range(limits.MAIL_PER_SOURCE + 1):
        client = Client(app.test_client())
        response = client.post("/sign-in", data={
            "email": "spoof{}@example.org".format(number)},
            headers={"CF-Connecting-IP": "10.0.0.{}".format(number)},
            environ_base={"REMOTE_ADDR": "192.0.2.7"})
        statuses.append(response.status_code)
    assert statuses[-1] == 429
    assert len(outbox.messages) == limits.MAIL_PER_SOURCE


# --- M1-C8: every path that sends mail has a finite bound --------------------

@pytest.fixture
def small_limits(monkeypatch):
    monkeypatch.setattr(limits, "SUPERVISOR_PER_ADDRESS", 2)
    monkeypatch.setattr(limits, "NOTICE_PER_ADDRESS", 1)


def test_requests_to_one_supervisor_are_bounded_and_roll_back(
        app, new_client, settings, outbox, small_limits):
    sent = []
    for number in range(3):
        student = new_client()
        as_student(student, "many{}@example.org".format(number))
        sent.append(submit_thesis(student, {
            "name1": "Име", "name2": "Бащино",
            "name3": "Фамилия" + "абв"[number]}).status_code)
    assert sent == [302, 302, 503]
    assert len(outbox.to("p.ivanov@example.org")) == 2
    third = record_for(settings, "many2@example.org")
    assert third["submitted_at"] is None
    connection = connect(settings)
    assert connection.execute("SELECT COUNT(*) FROM requests").fetchone()[0] \
        == 2
    assert connection.execute("SELECT COUNT(*) FROM attempts WHERE bucket = ?",
                              (limits.bucket("supervisor",
                                             "p.ivanov@example.org"),)
                              ).fetchone()[0] == 2
    connection.close()


def test_four_submissions_to_one_supervisor_leave_four_limit_records(
        app, new_client, settings, outbox):
    for number in range(4):
        student = new_client()
        as_student(student, "four{}@example.org".format(number))
        submit_thesis(student, {"name3": "Фамилия" + "абвг"[number]})
    connection = connect(settings)
    assert connection.execute("SELECT COUNT(*) FROM attempts WHERE bucket = ?",
                              (limits.bucket("supervisor",
                                             "p.ivanov@example.org"),)
                              ).fetchone()[0] == 4
    connection.close()


def test_a_bounded_replacement_rolls_back(app, client, new_client, settings,
                                          outbox, small_limits, staff):
    as_student(client, "replace@example.org")
    submit_thesis(client)
    record = record_for(settings, "replace@example.org")
    viewer = as_staff(new_client(), staff)
    base = "/staff/application/{}".format(record["id"])
    assert viewer.post(base + "/supervisor",
                       data={"professor": "dimitrov"}).status_code == 302
    assert viewer.post(base + "/supervisor",
                       data={"professor": "ivanov"}).status_code == 302
    # Two requests to ivanov were sent in this hour already.
    assert viewer.post(base + "/supervisor",
                       data={"professor": "dimitrov"}).status_code == 302
    response = viewer.post(base + "/supervisor",
                           data={"professor": "ivanov"})
    assert response.status_code == 503
    assert record_for(settings, "replace@example.org")["professor"] == \
        "dimitrov"


def test_a_bounded_decline_notice_keeps_the_answer_and_says_so(
        app, new_client, settings, outbox, small_limits):
    student = new_client()
    as_student(student, "notice2@example.org")
    submit_thesis(student)
    page = new_client().post(link_from(outbox.messages[0], "decline"),
                             data={}).get_data(as_text=True)
    assert "писмото до студента не" not in page
    student.post("/application/supervisor", data={"professor": "georgieva"})
    request = [m for m in outbox.messages
               if m.intended == "e.georgieva@example.org"][-1]
    page = new_client().post(link_from(request, "decline"),
                             data={}).get_data(as_text=True)
    assert "писмото до студента не" in page          # over the notice bound
    record = record_for(settings, "notice2@example.org")
    connection = connect(settings)
    assert applications.current_request(connection, record["id"])["answer"] \
        == "decline"
    connection.close()
    assert len(outbox.to("notice2@example.org")) == 1


def test_a_bounded_reminder_stays_due(app, client, settings, outbox, clock,
                                      monkeypatch):
    from superassistant import supervision
    as_student(client, "due@example.org")
    submit_thesis(client)
    monkeypatch.setattr(limits, "SUPERVISOR_PER_ADDRESS", 1)
    clock.advance(days=3, minutes=-30)                # the request's hour
    clock.advance(minutes=31)
    loaded = lists.load(settings.lists_dir)
    connection = connect(settings)
    monkeypatch.setattr(limits, "WINDOW", 4 * 86400)  # still counts it
    sent, failed = supervision.remind(connection, settings, loaded)
    assert sent == [] and len(failed) == 1
    assert connection.execute("SELECT reminders_sent FROM requests"
                              ).fetchone()[0] == 0
    monkeypatch.setattr(limits, "WINDOW", 3600)
    sent, failed = supervision.remind(connection, settings, loaded)
    assert len(sent) == 1 and failed == []
    connection.close()


def test_all_mail_together_is_bounded(app, new_client, settings, outbox,
                                      monkeypatch):
    monkeypatch.setattr(limits, "ALL_MAIL", 2)
    codes = []
    for number in range(3):
        client = Client(app.test_client())
        codes.append(client.post("/sign-in", data={
            "email": "all{}@example.org".format(number)},
            environ_base={"REMOTE_ADDR": "192.0.2.{}".format(number)}
        ).status_code)
    assert codes[:2] == [302, 302] and codes[2] == 429
    assert len(outbox.messages) == 2


# --- M1-C9: the report's header in the adopted size --------------------------

def test_the_report_header_is_times_new_roman_12_pt():
    from pathlib import Path
    template = Path(__file__).resolve().parent.parent / "templates" \
        / "doklad_temi.docx"
    header = Document(str(template)).tables[0].rows[0]
    runs = [r for c in header.cells for p in c.paragraphs for r in p.runs]
    assert runs and all(r.font.size is not None and r.font.size.pt == 12
                        for r in runs)
    assert all(r.font.name == "Times New Roman" for r in runs)


# --- Second round (Codex on bd50873) -------------------------------------------

def test_after_thesis_to_exam_staff_still_see_the_thesis_and_its_answers(
        client, new_client, settings, outbox, staff):
    as_student(client, "kept@example.org")
    submit_thesis(client)
    new_client().post(link_from(outbox.messages[0], "decline"),
                      data={"comment": "Нямам място този семестър."})
    record = record_for(settings, "kept@example.org")
    viewer = as_staff(new_client(), staff)
    base = "/staff/application/{}".format(record["id"])
    assert viewer.post(base + "/type",
                       data={"grad_type": "exam"}).status_code == 302
    page = viewer.get(base).get_data(as_text=True)
    connection = connect(settings)
    (request,) = applications.requests_of(connection, record["id"])
    connection.close()
    for shown in (THESIS_TAB_1["topic"], "Обосновка", "втори ред",
                  "Глава 2", "доц. д-р Петър Иванов", "p.ivanov@example.org",
                  "Не потвърждавам", "Нямам място този семестър.",
                  db.local_datetime(request["answered_at"]),
                  db.local_datetime(request["replaced_at"]),
                  "Запазени данни за теза"):
        assert shown in page, shown


@pytest.mark.parametrize("typed, expected", [
    ({"sup_name": "Анна Стефанова".replace("а", "a"), "sup_title": "проф. д-р",
      "sup_email": "a@example.net"}, "смесени латински и кирилски"),
    ({"sup_name": "Анна\x0bСтефанова", "sup_title": "проф. д-р",
      "sup_email": "a@example.net"}, "U+000B"),
    ({"professor": "nobody"}, "Изберете стойност от списъка"),
])
def test_a_refused_type_change_says_what_is_wrong(
        client, new_client, settings, outbox, staff, typed, expected):
    submit_exam(client, "detail{}@example.org".format(len(expected)))
    record = record_for(settings, "detail{}@example.org".format(
        len(expected)))
    viewer = as_staff(new_client(), staff)
    data = {"grad_type": "thesis"}
    data.update(typed)
    response = viewer.post("/staff/application/{}/type".format(record["id"]),
                           data=data)
    page = response.get_data(as_text=True)
    assert response.status_code == 400
    assert expected in page
    assert "За теза липсват" not in page


def test_only_genuinely_empty_fields_are_called_missing(
        client, new_client, settings, outbox, staff):
    submit_exam(client, "empty@example.org")
    record = record_for(settings, "empty@example.org")
    viewer = as_staff(new_client(), staff)
    page = viewer.post("/staff/application/{}/type".format(record["id"]),
                       data={"grad_type": "thesis", "professor": "ivanov"}
                       ).get_data(as_text=True)
    assert "За теза липсват" in page
    assert "Тема на тезата" in page and "смесени" not in page


def test_a_lists_folder_removed_after_installation_stops_the_start(tmp_path):
    import shutil
    settings = make_settings(tmp_path / "data")
    create_app(settings)
    shutil.rmtree(str(settings.lists_dir))
    with pytest.raises(StartupError) as raised:
        create_app(settings)
    assert "missing" in str(raised.value)
    assert not settings.lists_dir.exists()          # nothing recreated


def test_the_first_start_is_recorded_with_the_records(tmp_path):
    settings = make_settings(tmp_path / "data")
    create_app(settings)
    connection = connect(settings)
    assert connection.execute("SELECT key FROM installation").fetchall()[0][
        0] == "lists"
    connection.close()


def test_an_installation_from_schema_1_keeps_its_rows_and_its_lists(tmp_path):
    """Schema 2 adds the installation table; an installation made before it
    keeps every row, and its lists, edited or not, are not rewritten."""
    import sqlite3
    data_dir = tmp_path / "data"
    data_dir.mkdir()
    path = data_dir / "superassistant.sqlite3"
    connection = sqlite3.connect(str(path))
    connection.executescript(db.MIGRATIONS[0] + ";PRAGMA user_version = 1;")
    connection.execute(
        "INSERT INTO applications (email, email_key, created_at, updated_at,"
        " name1) VALUES ('old@example.org', 'old@example.org', 'x', 'x', "
        "'Старо')")
    connection.commit()
    connection.close()
    lists.write_starter(data_dir / "lists")
    edited = (data_dir / "lists" / lists.TITLES)
    edited.write_text("звание\nпроф. д-р\nдоц. д-р\nгл. ас. д-р\nас. д-р\n"
                      "акад.\n", encoding="utf-8")
    settings = make_settings(data_dir)
    create_app(settings)
    connection = connect(settings)
    assert connection.execute("SELECT name1 FROM applications").fetchone()[0] \
        == "Старо"
    assert connection.execute("PRAGMA user_version").fetchone()[0] == len(
        db.MIGRATIONS)
    connection.close()
    assert "акад." in edited.read_text(encoding="utf-8")


def test_the_signature_block_is_kept_together_in_a_filled_report(settings):
    from docx.text.paragraph import Paragraph
    from superassistant import documents
    loaded = lists.load(settings.lists_dir)
    rows = [documents.report_row(i + 1, {
        "id": i, "name1": "Име", "name2": "Бащино", "name3": "Ф" + str(i),
        "fn": "F{}".format(i), "topic": "Тема", "professor": "ivanov",
        "sup_name": "", "sup_title": "", "programme": "BSNADM01",
        "grad_type": "thesis"}, loaded) for i in range(3)]
    document = Document(io.BytesIO(documents.doklad_temi(
        "BSNADM01", rows, "2026-10-10T09:00:00+00:00", loaded)))
    body = list(document.element.body.iterchildren())
    after = [Paragraph(el, document)
             for el in body[body.index(document.tables[0]._tbl) + 1:]
             if el.tag.endswith("}p")]
    texts = [p.text for p in after]
    dated = max(i for i, text in enumerate(texts) if "Дата:" in text)
    block = after[:dated + 1]
    assert any("Приложение" in p.text for p in block)
    assert any("Виктор Аврамов/" in p.text for p in block)
    assert all(p.paragraph_format.keep_together for p in block)
    assert all(p.paragraph_format.keep_with_next for p in block[:-1])
    assert after[dated] is after[-1]                # no empty page-maker after
    header = document.tables[0].rows[0]
    assert {r.font.size.pt for c in header.cells for p in c.paragraphs
            for r in p.runs} == {12.0}


# --- Found by Opus 5.5's own check of 159925f --------------------------------

def schema_1_installation(data_dir):
    """A data directory as an installation before schema 2 left it."""
    import sqlite3
    data_dir.mkdir(parents=True)
    connection = sqlite3.connect(str(data_dir / "superassistant.sqlite3"))
    connection.executescript(db.MIGRATIONS[0] + ";PRAGMA user_version = 1;")
    connection.close()
    lists.write_starter(data_dir / "lists")


@pytest.mark.parametrize("remove", ["folder", lists.PROFESSORS])
def test_an_upgraded_installation_with_lists_missing_stops_the_start(
        tmp_path, remove):
    import shutil
    data_dir = tmp_path / "data"
    schema_1_installation(data_dir)
    if remove == "folder":
        shutil.rmtree(str(data_dir / "lists"))
    else:
        (data_dir / "lists" / remove).unlink()
    with pytest.raises(StartupError) as raised:
        create_app(make_settings(data_dir))
    assert "missing" in str(raised.value)
    if remove == "folder":
        assert not (data_dir / "lists").exists()
    else:
        assert not (data_dir / "lists" / remove).exists()


def test_an_upgraded_installation_is_recorded_as_installed(tmp_path):
    data_dir = tmp_path / "data"
    schema_1_installation(data_dir)
    settings = make_settings(data_dir)
    create_app(settings)
    connection = connect(settings)
    assert [r[0] for r in connection.execute("SELECT key FROM installation")] \
        == ["lists"]
    connection.close()


def test_a_failed_first_start_leaves_no_partial_list(tmp_path, monkeypatch):
    import os
    data_dir = tmp_path / "data"
    settings = make_settings(data_dir)
    real_replace = os.replace
    calls = []

    def failing(src, dst):
        calls.append(dst)
        if len(calls) == 2:
            raise OSError(28, "No space left on device")
        real_replace(src, dst)
    monkeypatch.setattr(lists.os, "replace", failing)
    with pytest.raises(OSError):
        create_app(settings)
    monkeypatch.setattr(lists.os, "replace", real_replace)
    create_app(settings)                        # the retry completes the lists
    loaded = lists.load(settings.lists_dir)
    assert "BSNADM01" in loaded.programmes and loaded.titles


TYPED_BAD_ADDRESS = {"professor": "", "sup_title": "проф. д-р",
                     "sup_name": "Анна Стефанова", "sup_email": "anna@example"}


def test_a_typed_address_that_is_not_one_is_wrong_not_missing(
        client, new_client, settings, outbox, staff):
    from superassistant import texts
    as_student(client, "badaddress@example.org")
    response = submit_thesis(client, TYPED_BAD_ADDRESS)
    page = response.get_data(as_text=True)
    assert response.status_code == 400
    assert texts.MAIL_BAD_ADDRESS in page
    assert texts.REFUSED_REQUIRED not in page.split('id="sup_email-error"')[
        1][:200]
    submit_exam(new_client(), "badtype@example.org")
    record = record_for(settings, "badtype@example.org")
    viewer = as_staff(new_client(), staff)
    base = "/staff/application/{}".format(record["id"])
    for field, value in thesis_data().items():
        viewer.post(base + "/correct", data={"field": field, "value": value})
    data = {"grad_type": "thesis"}
    data.update(TYPED_BAD_ADDRESS)
    page = viewer.post(base + "/type", data=data).get_data(as_text=True)
    assert texts.MAIL_BAD_ADDRESS in page and "За теза липсват" not in page


def test_a_retained_title_alone_is_shown(client, new_client, settings, staff):
    as_student(client, "titleonly@example.org")
    save(client, 1, {"name1": "Мила", "name2": "Петрова", "name3": "Нейкова",
                     "fn": "700001", "grad_type": "thesis",
                     "sup_title": "проф. д-р"})
    save(client, 1, {"name1": "Мила", "name2": "Петрова", "name3": "Нейкова",
                     "fn": "700001", "grad_type": "exam",
                     "credits_declared": "1"}, "submit")
    record = record_for(settings, "titleonly@example.org")
    assert record["sup_title"] == "проф. д-р"
    page = as_staff(new_client(), staff).get(
        "/staff/application/{}".format(record["id"])).get_data(as_text=True)
    assert "Запазени данни за теза" in page
    assert "проф. д-р" in page.split('id="thesis-fields"')[1].split("</dl>")[0]
