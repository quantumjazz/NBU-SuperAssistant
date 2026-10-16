"""Each applicant's forms in a real browser (constitution 5.5; the owner, 16
October 2026): the student, at phone width, reads „Попълнени документи“ and
the line saying there is nothing to print or hand in; staff, at desktop
width, download that student's заявление-декларация and Приложение 2 from
the application's page, and the заявление alone for a state exam.
Everything is invented.
"""

import io

import pytest
from docx import Document

from superassistant import texts
from tests.browser.conftest import DESKTOP, PHONE
from tests.browser.test_application_browser import STUDENT, submit_thesis
from tests.conftest import insert_application, make_account

sync_api = pytest.importorskip("playwright.sync_api")
expect = sync_api.expect

pytestmark = pytest.mark.browser

NAME = " ".join(STUDENT[:3])


def downloaded_text(download):
    document = Document(io.BytesIO(open(download.path(), "rb").read()))
    text = [p.text for p in document.paragraphs]
    text += [c.text for t in document.tables for r in t.rows for c in r.cells]
    return "\n".join(text)


def test_the_student_reads_the_line_and_staff_download_the_forms(
        live_server, open_page):
    settings = live_server.settings
    make_account(settings, "staff6", password="служител-6")
    student = open_page(PHONE)
    submit_thesis(student, live_server, "forms@example.org", STUDENT)
    expect(student.locator("#forms h2")).to_have_text(texts.FORMS_TITLE)
    expect(student.locator("#forms-lead")).to_have_text(texts.FORMS_LEAD)
    expect(student.locator("#forms-lead")).to_be_visible()
    exam = insert_application(settings, "exam6@example.org",
                              grad_type="exam", name1="Лора",
                              name2="Методиева", name3="Ненова",
                              fn="F200002")

    desk = open_page(DESKTOP)
    desk.goto(live_server.url + "/staff/sign-in")
    desk.fill("#login", "staff6")
    desk.fill("#password", "служител-6")
    desk.get_by_role("button", name=texts.STAFF_SIGNIN_BUTTON).click()
    desk.get_by_role("link", name=NAME).click()
    card = desk.locator("#staff-forms")
    expect(card.locator("h2")).to_have_text(texts.STAFF_FORMS_TITLE)
    for label, expected in ((texts.FORMS_ZAIAVLENIE, NAME),
                            (texts.FORMS_PRILOZHENIE2, STUDENT[4])):
        with desk.expect_download() as download:
            card.get_by_role("link", name=label).click()
        assert download.value.suggested_filename.endswith(
            "-F" + STUDENT[3] + ".docx")
        assert expected in downloaded_text(download.value)

    desk.goto(live_server.url + "/staff/application/{}".format(exam))
    card = desk.locator("#staff-forms")
    expect(card.get_by_role("link")).to_have_count(1)
    with desk.expect_download() as download:
        card.get_by_role("link", name=texts.FORMS_ZAIAVLENIE).click()
    assert download.value.suggested_filename == \
        "zaiavlenie-deklaracia-{}-F200002.docx".format(exam)
    assert "Лора Методиева Ненова" in downloaded_text(download.value)
