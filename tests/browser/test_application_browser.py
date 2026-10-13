"""The application in a real browser (NOW.md proof 5).

At phone width (390 px): a thesis application from the sign-in link's
„Влез“ page to download, with the first tab's reminder shown, the professors
and programmes offered read from the list files and no field asked twice; a
state-exam application with no second tab, its page then showing
"submitted"; a supervisor's answer through each button, with a comment, the
answer then shown to the student and in the staff table. At desktop width:
the staff table, one opened application and a check recorded.
"""

import io
import re

import pytest
from docx import Document

from superassistant import lists, texts
from tests.browser.conftest import DESKTOP, PHONE
from tests.conftest import make_account

sync_api = pytest.importorskip("playwright.sync_api")
expect = sync_api.expect

pytestmark = pytest.mark.browser


def sign_in_by_link(page, server, address):
    """Ask for the email, then open its link and press „Влез“."""
    page.goto(server.url + "/")
    page.fill("#email", address)
    page.get_by_role("button", name=texts.SIGNIN_SEND).click()
    message = server.latest(address)
    link = re.search(r"https?://\S+/link/\S+", message.text).group(0)
    page.goto(link)
    # Opening the link only shows the „Влез“ page.
    expect(page.get_by_role("button", name=texts.SIGNIN_LINK_BUTTON)) \
        .to_be_visible()
    page.get_by_role("button", name=texts.SIGNIN_LINK_BUTTON).click()
    expect(page).to_have_url(server.url + "/application")


def asked(page):
    """The names of the fields a page asks, apart from the hidden ones."""
    return set(page.eval_on_selector_all(
        "form#application-form input:not([type=hidden]), "
        "form#application-form textarea, form#application-form select",
        "els => els.map(e => e.name)"))


def fill_thesis_tab_1(page, student):
    page.fill("#name1", student[0])
    page.fill("#name2", student[1])
    page.fill("#name3", student[2])
    page.fill("#fn", student[3])
    page.check("input[name=grad_type][value=thesis]")
    page.fill("#phone", "0888 123 456")
    page.select_option("#programme", "MSNSMQ09")
    page.fill("#topic", student[4])
    page.fill("#average_grade", "5,75")
    page.select_option("#professor", "georgieva")
    page.check("#credits_declared")


def fill_thesis_tab_2(page):
    for field in ("rationale", "aims", "object_subject", "methodology",
                  "information_base", "plan", "schedule"):
        page.fill("#" + field, "Раздел " + field + "\nвтори ред")


def submit_thesis(page, server, address, student):
    sign_in_by_link(page, server, address)
    fill_thesis_tab_1(page, student)
    page.locator("#next-button").click()
    expect(page).to_have_url(server.url + "/application?tab=2")
    fill_thesis_tab_2(page)
    page.locator("#submit-thesis").click()
    expect(page.locator("#status")).to_be_visible()


STUDENT = ("Десислава", "Петкова", "Маринова", "201234",
           "Качество на услугите в публичния сектор")


def test_a_thesis_application_from_the_link_to_the_download(live_server,
                                                             open_page):
    page = open_page(PHONE)
    sign_in_by_link(page, live_server, "desislava@example.org")
    expect(page.locator("#advice")).to_have_text(texts.ADVICE)

    # What is offered comes from the list files.
    loaded = lists.load(live_server.settings.lists_dir)
    offered = page.eval_on_selector_all("#professor option",
                                        "els => els.map(e => e.textContent)")
    for code, professor in loaded.professors.items():
        assert "{} {}".format(professor["title"], professor["name"]) in offered
    offered = page.eval_on_selector_all("#programme option",
                                        "els => els.map(e => e.value)")
    assert set(loaded.programmes) <= set(offered)

    fill_thesis_tab_1(page, STUDENT)
    first = asked(page)
    page.locator("#next-button").click()
    fill_thesis_tab_2(page)
    second = asked(page)
    assert first and second and not (first & second)    # none asked twice
    page.locator("#submit-thesis").click()
    expect(page.locator("#status")).to_contain_text(texts.STATUS_AWAITED)

    for link, expected in (("Заявление-декларация", "Десислава Петкова "
                            "Маринова"),
                           ("Приложение 2", "Качество на услугите")):
        with page.expect_download() as download:
            page.get_by_role("link", name=re.compile(link)).click()
        body = open(download.value.path(), "rb").read()
        text = "\n".join(p.text for p in Document(io.BytesIO(body)).paragraphs)
        text += "\n".join(c.text for t in Document(io.BytesIO(body)).tables
                          for r in t.rows for c in r.cells)
        assert expected in text
    # The request went to the supervisor from the list.
    assert live_server.latest("e.georgieva@example.org") is not None


def test_a_state_exam_application_has_no_second_tab(live_server, open_page):
    page = open_page(PHONE)
    sign_in_by_link(page, live_server, "exam@example.org")
    page.fill("#name1", "Калоян")
    page.fill("#name2", "Иванов")
    page.fill("#name3", "Стоев")
    page.fill("#fn", "300123")
    page.check("input[name=grad_type][value=exam]")
    page.check("#credits_declared")
    expect(page.locator("#tab2-link")).to_be_hidden()
    expect(page.locator("#thesis-fields")).to_be_hidden()
    page.locator("#submit-exam").click()
    expect(page.locator("#status")).to_contain_text("Подадено на")
    expect(page.locator("#tabs")).to_have_count(0)
    assert live_server.outbox.to("e.georgieva@example.org") == []


def test_a_supervisor_answers_through_each_button(live_server, open_page):
    students = {
        "first@example.org": ("Анелия", "Рускова", "Борисова", "400001",
                              "Тема първа"),
        "second@example.org": ("Борислав", "Георгиев", "Христов", "400002",
                               "Тема втора"),
    }
    pages = {}
    for address, student in students.items():
        pages[address] = open_page(PHONE)
        submit_thesis(pages[address], live_server, address, student)
    requests = live_server.outbox.to("e.georgieva@example.org")
    assert len(requests) == 2
    answers = {"first@example.org": "confirm", "second@example.org": "decline"}
    for message, (address, choice) in zip(requests, answers.items()):
        assert students[address][0] in message.text
        link = re.search(r"https?://\S+/answer/\S+/" + choice,
                         message.text).group(0)
        supervisor = open_page(PHONE)
        supervisor.goto(link)
        supervisor.fill("#comment", "Коментар за " + choice)
        supervisor.locator("#record-answer").click()
        expect(supervisor.locator("#answer-state")).to_contain_text(
            texts.ANSWER_WORD[choice])
    pages["first@example.org"].reload()
    expect(pages["first@example.org"].locator("#status")).to_contain_text(
        texts.STATUS_CONFIRMED)
    pages["second@example.org"].reload()
    expect(pages["second@example.org"].locator("#status")).to_contain_text(
        texts.STATUS_DECLINED)
    expect(pages["second@example.org"].locator("#status")).not_to_contain_text(
        "Коментар")

    staff = make_account(live_server.settings, "staff2",
                         password="служител-парола-2")
    desk = open_page(DESKTOP)
    desk.goto(live_server.url + "/staff/sign-in")
    desk.fill("#login", "staff2")
    desk.fill("#password", "служител-парола-2")
    desk.get_by_role("button", name=texts.STAFF_SIGNIN_BUTTON).click()
    table = desk.locator("#applications")
    expect(table).to_contain_text(texts.STAFF_ANSWER["confirm"])
    expect(table).to_contain_text(texts.STAFF_ANSWER["decline"])
    assert staff is not None


def test_staff_table_one_application_and_a_check(live_server, open_page):
    student_page = open_page(PHONE)
    submit_thesis(student_page, live_server, "checked@example.org", STUDENT)
    draft = open_page(PHONE)
    sign_in_by_link(draft, live_server, "draft@example.org")
    draft.fill("#name1", "Недовършена")
    draft.locator("button[value=save]").click()

    make_account(live_server.settings, "staff3", password="служител-3")
    desk = open_page(DESKTOP)
    desk.goto(live_server.url + "/staff/sign-in")
    desk.fill("#login", "staff3@example.org")
    desk.fill("#password", "служител-3")
    desk.get_by_role("button", name=texts.STAFF_SIGNIN_BUTTON).click()
    table = desk.locator("#applications")
    expect(table).to_contain_text("Десислава Петкова Маринова")
    expect(table).not_to_contain_text("Недовършена")
    for column in texts.STAFF_COLUMNS:
        expect(table.locator("thead")).to_contain_text(column)
    row = table.locator("tbody tr").first
    expect(row).to_have_class("waiting")
    expect(row).to_contain_text(texts.STAFF_WAITING_SUPERVISOR)
    expect(row).to_contain_text(texts.STAFF_WAITING_CHECK)

    table.get_by_role("link", name="Десислава Петкова Маринова").click()
    fields = desk.locator("main")
    for value in ("F201234", "0888 123 456", "5,75", STUDENT[4],
                  "Раздел plan"):
        expect(fields).to_contain_text(value)
    desk.check("input[name=result][value=passed]")
    desk.fill("#note", "Данните съвпадат.")
    desk.locator("#record-check").click()
    expect(desk.locator("#checks")).to_contain_text("Данните съвпадат.")
    expect(desk.locator("#checks")).to_contain_text("staff3")
    desk.goto(live_server.url + "/staff")
    expect(desk.locator("#applications tbody tr").first).not_to_contain_text(
        texts.STAFF_WAITING_CHECK)


def test_a_refused_thesis_entry_keeps_the_thesis_choice_and_shows_its_error(
        live_server, open_page):
    """Codex M1-C1: a refusal redisplayed the page with no graduation type
    chosen, the error hidden inside the thesis fields."""
    page = open_page(PHONE)
    sign_in_by_link(page, live_server, "refused-topic@example.org")
    fill_thesis_tab_1(page, STUDENT[:4] + ("Тeма с латинско e",))
    page.locator("#next-button").click()
    expect(page.locator("input[name=grad_type][value=thesis]")).to_be_checked()
    expect(page.locator("#thesis-fields")).to_be_visible()
    expect(page.locator("#topic-error")).to_be_visible()
    expect(page.locator("#topic-error")).to_contain_text("e")
    expect(page.locator("#tab2-link")).to_be_visible()


def test_switching_to_a_state_exam_ignores_the_hidden_thesis_fields(
        live_server, open_page):
    """Codex M1-C2: an invalid hidden thesis entry blocked a state exam; and
    thesis entries saved earlier are kept, not deleted."""
    page = open_page(PHONE)
    sign_in_by_link(page, live_server, "switch@example.org")
    fill_thesis_tab_1(page, STUDENT)
    page.locator("button[value=save]").click()
    expect(page.locator(".flash.ok")).to_be_visible()
    page.fill("#topic", "Тeма с латинско e")          # invalid, then hidden
    page.check("input[name=grad_type][value=exam]")
    expect(page.locator("#thesis-fields")).to_be_hidden()
    page.locator("#submit-exam").click()
    expect(page.locator("#status")).to_contain_text("Подадено на")
    from superassistant import applications, db
    connection = db.connect(live_server.settings.db_path)
    try:
        record = applications.by_email_key(connection, "switch@example.org")
    finally:
        connection.close()
    assert record["grad_type"] == "exam" and record["submitted_at"]
    assert record["topic"] == STUDENT[4]              # kept as saved
