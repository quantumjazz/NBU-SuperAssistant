"""F№ typed as digits after a fixed „F“ (the task of 13 October 2026;
constitution 5.3): the student's form, staff's correction, the values stored
before the rule, and the documents, which print the stored value as before.
The rule itself is tested over generated text in tests/test_entry_rules.py.
Everyone is invented.
"""

import io
import re

import docx

from superassistant import applications, texts
from tests.conftest import connect, insert_application, make_account
from tests.test_application_flow import (THESIS_TAB_1, as_staff, as_student,
                                         record_for, save, submit_thesis)
from tests.test_documents import every_paragraph


def test_the_field_shows_a_fixed_letter_and_takes_digits(client, settings):
    as_student(client, "digits@example.org")
    page = client.get("/application").get_data(as_text=True)
    assert re.search(r'<div class="prefixed"><span class="prefix" '
                     r'aria-hidden="true">F</span><input type="text" id="fn" '
                     r'name="fn" inputmode="numeric"', page)
    save(client, 1, dict(THESIS_TAB_1, fn="104999"))
    assert record_for(settings, "digits@example.org")["fn"] == "F104999"
    page = client.get("/application").get_data(as_text=True)
    assert re.search(r'id="fn" [^>]*value="104999"', page)


def test_anything_but_digits_is_refused_and_shown_as_typed(client, settings):
    as_student(client, "letters@example.org")
    save(client, 1, dict(THESIS_TAB_1, fn="104999"))
    for typed in ("F104999", "Ф104999", "104 999", "f104999"):
        response = save(client, 1, dict(THESIS_TAB_1, fn=typed))
        page = response.get_data(as_text=True)
        assert response.status_code == 400
        assert re.search(r'id="fn-error">' + re.escape(texts.REFUSED_FN),
                         page)
        assert re.search(r'id="fn" [^>]*value="{}"'.format(typed), page)
        assert record_for(settings, "letters@example.org")["fn"] == \
            "F104999"


def test_a_draft_stored_before_the_rule_shows_its_value_to_retype(
        client, settings):
    insert_application(settings, "older@example.org", submitted=False,
                       fn="104999", grad_type="exam")
    as_student(client, "older@example.org")
    page = client.get("/application").get_data(as_text=True)
    assert re.search(r'id="fn" [^>]*value="104999"', page)
    save(client, 1, {"name1": "Ана", "fn": "104999", "grad_type": "exam"})
    assert record_for(settings, "older@example.org")["fn"] == "F104999"


SHAPES = {"rule": "F123456", "digits": "123456", "other": "Ф-12/3"}


def test_the_stored_shapes_are_shown_to_staff_as_stored_and_corrected(
        new_client, settings):
    """Nothing stored is rewritten: each value shows as it is, and staff
    correct it under the rule, with or without the letter."""
    numbers = dict((shape, insert_application(
        settings, shape + "@example.org", grad_type="exam", name1="Ана",
        name2="Бойкова", name3=shape, fn=value))
        for shape, value in SHAPES.items())
    viewer = as_staff(new_client(), make_account(settings, "staff1"))
    table = viewer.get("/staff").get_data(as_text=True)
    for shape, number in numbers.items():
        row = table.split('data-id="{}"'.format(number))[1].split("</tr>")[0]
        assert "<td>{}</td>".format(SHAPES[shape]) in row
        page = viewer.get("/staff/application/{}".format(number)).get_data(
            as_text=True)
        assert "<dd>{}</dd>".format(SHAPES[shape]) in page
    for shape, typed, stored in (("digits", "123456", "F123456"),
                                 ("other", "f777", "F777"),
                                 ("rule", "F123457", "F123457")):
        base = "/staff/application/{}".format(numbers[shape])
        refused = viewer.post(base + "/correct", data={"field": "fn",
                                                       "value": "Ф777"})
        assert refused.status_code == 400
        assert texts.REFUSED_FN in refused.get_data(as_text=True)
        assert viewer.post(base + "/correct", data={
            "field": "fn", "value": typed}).status_code == 302
        connection = connect(settings)
        (correction,) = applications.corrections_of(connection,
                                                    numbers[shape])
        connection.close()
        assert (correction["old_value"], correction["new_value"]) == (
            SHAPES[shape], stored)


def test_the_forms_and_the_report_print_the_prefixed_value(
        client, new_client, settings, outbox):
    from superassistant import lists, reports
    from tests.test_application_flow import link_from
    as_student(client, "printed@example.org")
    submit_thesis(client, {"fn": "090737"})
    assert "F090737" in outbox.messages[0].text
    for form in ("zaiavlenie", "prilozhenie2"):
        reply = client.get("/application/forms/" + form + ".docx")
        text = "\n".join(p.text for p in every_paragraph(
            docx.Document(io.BytesIO(reply.data))))
        assert "F090737" in text and "FF090737" not in text
    new_client().post(link_from(outbox.messages[0], "confirm"), data={})
    record = record_for(settings, "printed@example.org")
    staff = make_account(settings, "staff1")
    connection = connect(settings)
    applications.record_check(connection, record["id"], "passed", "",
                              staff["id"])
    report_id = reports.generate(connection, settings,
                                 lists.load(settings.lists_dir),
                                 record["programme"], staff["id"])
    path = reports.path(settings, reports.get(connection, report_id))
    connection.close()
    table = docx.Document(str(path)).tables[0]
    assert [c.text for c in table.rows[1].cells][1] == "F090737"
