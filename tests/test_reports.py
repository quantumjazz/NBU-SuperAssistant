"""The topic-approval report's rules (NOW.md requirement 7; proof 1): which
faculty and dean a programme leads to, which applications a report lists and
leaves out, that earlier reports are kept with their date and students, and
that a programme with nothing to report makes no report. Over generated list
files and generated applications; everyone is invented.
"""

import csv
import hashlib
import io
import os
import secrets
import shutil
import sqlite3
import string
import tempfile
from pathlib import Path

import docx
import pytest
from hypothesis import event, given, settings as hsettings, strategies as st

from superassistant import db, documents, lists, reports, texts
from tests.conftest import (Client, insert_application, insert_check,
                            insert_request, make_account, make_settings,
                            prepare_data_dir)
from tests.test_documents import line


class Data(object):
    """A fresh data directory for one generated example."""

    def __enter__(self):
        self.root = Path(tempfile.mkdtemp())
        prepare_data_dir(self.root / "data")
        self.settings = make_settings(self.root / "data")
        self.lists = lists.load(self.settings.lists_dir)
        self.connection = db.connect(self.settings.db_path)
        self.account = add_account(self.connection)
        return self

    def __exit__(self, *exc):
        self.connection.close()
        shutil.rmtree(str(self.root), ignore_errors=True)


def add_account(connection):
    """A staff account written straight into the schema; no slow hashing,
    since these tests never sign in with it."""
    stamp = db.now()
    name = "staff" + secrets.token_hex(4)
    cursor = connection.execute(
        "INSERT INTO accounts (username, username_key, email, email_key, "
        "algorithm, salt, iterations, dklen, hash, session_tag, created_at, "
        "approved_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
        (name, name, name + "@example.org", name + "@example.org", "none",
         "00", 1, 1, "00", secrets.token_hex(8), stamp, stamp))
    return cursor.lastrowid


def write_csv(path, header, rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(header)
    for row in rows:
        writer.writerow(row)
    path.write_text(buffer.getvalue(), encoding="utf-8")


def cell_text():
    """A list entry: one line, no space at either end, any characters."""
    return line(max_size=30)


# --- Which faculty and dean a programme leads to -----------------------------

@st.composite
def list_files(draw):
    titles = draw(st.lists(cell_text(), min_size=1, max_size=4, unique=True))
    codes = draw(st.lists(st.text(string.ascii_letters + string.digits
                                  + "АБВГДЕ-", min_size=1, max_size=6),
                          min_size=1, max_size=4, unique=True))
    faculties = [[code, draw(cell_text()), draw(st.sampled_from(titles)),
                  draw(cell_text())] for code in codes]
    programme_codes = draw(st.lists(
        st.text(string.ascii_letters + "ЖЗИК-", min_size=1, max_size=6),
        min_size=1, max_size=6, unique=True))
    programmes = [[code, draw(st.sampled_from(lists.DEGREES)),
                   draw(cell_text()), draw(st.sampled_from(lists.FORMS)),
                   draw(st.sampled_from(codes))] for code in programme_codes]
    return titles, faculties, programmes


@hsettings(max_examples=50, deadline=None)
@given(list_files(), st.data())
def test_a_programme_leads_to_its_faculty_and_dean(files, data):
    titles, faculties, programmes = files
    with Data() as d:
        lists_dir = d.settings.lists_dir
        write_csv(lists_dir / lists.TITLES, lists.HEADERS[lists.TITLES],
                  [[t] for t in titles])
        write_csv(lists_dir / lists.FACULTIES,
                  lists.HEADERS[lists.FACULTIES], faculties)
        write_csv(lists_dir / lists.PROGRAMMES,
                  lists.HEADERS[lists.PROGRAMMES], programmes)
        write_csv(lists_dir / lists.PROFESSORS,
                  lists.HEADERS[lists.PROFESSORS],
                  [["prof", titles[0], "Петър Иванов", "p@example.org"]])
        loaded = lists.load(lists_dir)
        by_code = {row[0]: row for row in faculties}
        for code, degree, name, form, faculty in programmes:
            expected = by_code[faculty]
            assert loaded.programme_faculty(code) == {
                "name": expected[1], "dean_title": expected[2],
                "dean": expected[3]}
        assert loaded.programme_faculty("no such programme") is None

        # And the report for one of them is addressed to that dean.
        code, degree, name, form, faculty = data.draw(
            st.sampled_from(programmes))
        number = insert_application(
            d.settings, "a@example.org", grad_type="thesis",
            programme=code, name1="Ана", name2="Бойкова", name3="Вълчева",
            fn="F1", topic="Тема", professor="prof")
        insert_request(d.settings, number, answer="confirm")
        insert_check(d.settings, number, "passed", d.account)
        report_id = reports.generate(d.connection, d.settings, loaded, code,
                                     d.account)
        report = reports.get(d.connection, report_id)
        assert report["faculty"] == faculty
        body = [p.text for p in docx.Document(
            str(reports.path(d.settings, report))).paragraphs]
        _, faculty_name, dean_title, dean = by_code[faculty]
        assert body[1] == (dean_title + " " + dean).upper()
        assert body[2] == "ДЕКАН НА " + faculty_name.upper()


# --- Which applications a report lists ---------------------------------------

PROGRAMMES = ("BSNADM01", "MSNMBA01")
NAMES = ["Ана", "Боян", "Вера", "Галин", "Дора", "Емил", "Жана", "Зорница",
         "Ивайло", "Йоана", "Калин", "Лия"]


@st.composite
def generated_application(draw):
    return {
        "programme": draw(st.sampled_from(PROGRAMMES)),
        # Weighted towards what qualifies, so both outcomes come up often.
        "grad_type": draw(st.sampled_from(["thesis"] * 3 + ["exam"])),
        "submitted": draw(st.sampled_from([True] * 3 + [False])),
        "name": draw(st.sampled_from(NAMES)) + " " + draw(cell_text()),
        "fn": "F" + draw(st.text(string.digits, min_size=1, max_size=6)),
        "requests": draw(st.lists(st.tuples(
            st.sampled_from(["confirm", "confirm", None, "decline"]),
            st.sampled_from([False, False, True])), max_size=3)),
        "checks": draw(st.lists(st.sampled_from(["passed", "passed",
                                                 "not_passed"]),
                                max_size=3)),
    }


def belongs(generated, programme):
    """The rule, written again from NOW.md requirement 7."""
    current = [answer for answer, replaced in generated["requests"]
               if not replaced]
    return (generated["programme"] == programme
            and generated["grad_type"] == "thesis"
            and generated["submitted"]
            and bool(current) and current[-1] == "confirm"
            and bool(generated["checks"])
            and generated["checks"][-1] == "passed")


def insert(d, generated, index):
    first, last = generated["name"].split(" ", 1)
    number = insert_application(
        d.settings, "s{}@example.org".format(index),
        submitted=generated["submitted"], grad_type=generated["grad_type"],
        programme=generated["programme"], name1=first, name2="Тестова",
        name3=last, fn=generated["fn"], topic="Тема " + str(index),
        professor="ivanov")
    for answer, replaced in generated["requests"]:
        insert_request(d.settings, number, answer=answer, replaced=replaced)
    for result in generated["checks"]:
        insert_check(d.settings, number, result, d.account)
    return number


@hsettings(max_examples=60, deadline=None)
@given(st.lists(generated_application(), max_size=8),
       st.sampled_from(PROGRAMMES))
def test_a_report_lists_exactly_the_confirmed_and_passed(applications,
                                                         programme):
    with Data() as d:
        numbers = [insert(d, generated, index)
                   for index, generated in enumerate(applications)]
        expected = {number for number, generated in zip(numbers,
                                                        applications)
                    if belongs(generated, programme)}
        event("listed" if expected else "nothing to report")
        listed = reports.qualifying(d.connection, programme)
        assert {row["id"] for row in listed} == expected
        assert reports.counts(d.connection).get(programme, 0) == len(expected)

        if not expected:
            with pytest.raises(reports.NothingToReport):
                reports.generate(d.connection, d.settings, d.lists,
                                 programme, d.account)
            assert reports.kept(d.connection) == []
            assert list(d.settings.reports_dir.iterdir()) == []
            return

        report_id = reports.generate(d.connection, d.settings, d.lists,
                                     programme, d.account)
        kept = reports.students(d.connection, report_id)
        assert [row["application_id"] for row in kept] == [
            row["id"] for row in listed]
        assert [row["position"] for row in kept] == list(
            range(1, len(expected) + 1))
        # In alphabetical order of the names as printed, then by F№.
        keys = [(row["name"].casefold(), row["fn"].casefold())
                for row in kept]
        assert keys == sorted(keys)
        table = docx.Document(str(reports.path(
            d.settings, reports.get(d.connection, report_id)))).tables[0]
        assert [[c.text for c in r.cells] for r in table.rows[1:]] == [
            [str(row["position"]), row["fn"], row["name"], row["topic"],
             row["supervisor"], ""] for row in kept]


# --- Earlier reports are kept ----------------------------------------------

def fingerprint(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


@hsettings(max_examples=25, deadline=None)
@given(st.lists(st.lists(generated_application(), min_size=1, max_size=4),
                min_size=2, max_size=4))
def test_earlier_reports_are_kept_with_their_date_and_students(rounds):
    clock = [db.parse("2026-10-12T09:00:00+00:00")]
    db.set_clock(lambda: clock[0])
    try:
        with Data() as d:
            made = []
            index = 0
            for round_ in rounds:
                for generated in round_:
                    insert(d, generated, index)
                    index += 1
                # A check that now fails, for one already reported.
                if made and made[-1]["students"]:
                    insert_check(d.settings,
                                 made[-1]["students"][0]["application_id"],
                                 "not_passed", d.account)
                clock[0] = clock[0].replace(day=clock[0].day + 1)
                try:
                    report_id = reports.generate(d.connection, d.settings,
                                                 d.lists, "BSNADM01", d.account)
                except reports.NothingToReport:
                    continue
                report = reports.get(d.connection, report_id)
                made.append({
                    "id": report_id, "row": dict(report),
                    "students": [dict(r) for r in reports.students(
                        d.connection, report_id)],
                    "file": fingerprint(reports.path(d.settings, report))})
            kept = reports.kept(d.connection)
            assert [r["id"] for r in kept] == [m["id"] for m in
                                               reversed(made)]
            for earlier in made:
                report = reports.get(d.connection, earlier["id"])
                assert dict(report) == earlier["row"]
                assert [dict(r) for r in reports.students(
                    d.connection, earlier["id"])] == earlier["students"]
                assert fingerprint(reports.path(d.settings, report)) == \
                    earlier["file"]
            assert len({m["row"]["filename"] for m in made}) == len(made)
            assert len(list(d.settings.reports_dir.iterdir())) == len(made)
            dates = [m["row"]["generated_at"] for m in made]
            assert dates == sorted(dates) and len(set(dates)) == len(dates)
    finally:
        db.set_clock(None)


def test_nothing_qualifying_makes_no_report(tmp_path):
    data_dir = tmp_path / "data"
    prepare_data_dir(data_dir)
    settings = make_settings(data_dir)
    connection = db.connect(settings.db_path)
    try:
        account = add_account(connection)
        number = insert_application(settings, "a@example.org",
                                    grad_type="thesis", programme="BSNADM01",
                                    name1="Ана", fn="F1", professor="ivanov")
        insert_request(settings, number, answer="confirm")
        insert_check(settings, number, "not_passed", account)
        with pytest.raises(reports.NothingToReport):
            reports.generate(connection, settings, lists.load(
                settings.lists_dir), "BSNADM01", account)
        assert reports.kept(connection) == []
        assert not settings.reports_dir.exists() or not list(
            settings.reports_dir.iterdir())
    finally:
        connection.close()


# --- The staff page --------------------------------------------------------

def staff_client(app, record):
    client = Client(app.test_client())
    with client.raw.session_transaction() as session:
        session["staff"] = record["id"]
        session["tag"] = record["session_tag"]
    return client


def test_staff_generate_list_and_download_reports(settings, app, clock):
    staff = make_account(settings)
    client = staff_client(app, staff)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", name2="Бойкова",
                                name3="Вълчева", fn="F300001",
                                topic="Тема", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])

    page = client.get("/staff/reports").get_data(as_text=True)
    assert documents.programme_label("BSNADM01", app.extensions[
        "superassistant"]["lists"]) in page
    assert "заявления за доклада: 1" in page
    assert reports_count(settings) == 0

    reply = client.post("/staff/reports", {"programme": "MSNMBA01"},
                        follow_redirects=True)
    assert "Няма заявление за доклада" in reply.get_data(as_text=True)
    assert reports_count(settings) == 0

    reply = client.post("/staff/reports", {"programme": "BSNADM01"},
                        follow_redirects=True)
    page = reply.get_data(as_text=True)
    assert "Докладът е създаден." in page
    assert "12.10.2026" in page and "/staff/reports/1.docx" in page
    reply = client.get("/staff/reports/1.docx")
    assert reply.status_code == 200
    assert reply.headers["Content-Disposition"].isascii()
    table = docx.Document(io.BytesIO(reply.data)).tables[0]
    assert [c.text for c in table.rows[1].cells][:3] == [
        "1", "F300001", "Ана Бойкова Вълчева"]

    clock.advance(days=1)
    client.post("/staff/reports", {"programme": "BSNADM01"})
    assert reports_count(settings) == 2
    assert client.get("/staff/reports/1.docx").data == reply.data
    assert client.get("/staff/reports/3.docx").status_code == 404
    assert client.post("/staff/reports", {"programme": "nope"}).status_code \
        == 302
    assert reports_count(settings) == 2


def reports_count(settings):
    connection = db.connect(settings.db_path)
    try:
        return connection.execute("SELECT COUNT(*) FROM reports").fetchone()[0]
    finally:
        connection.close()


def test_only_approved_staff_reach_reports(settings, app):
    waiting = make_account(settings, username="waiting", approved=False)
    student = Client(app.test_client())
    with student.raw.session_transaction() as session:
        session["student"] = "a@example.org"
    for client in (Client(app.test_client()), student,
                   staff_client(app, waiting)):
        for url in ("/staff/reports", "/staff/reports/1.docx"):
            reply = client.get(url)
            assert reply.status_code == 302
            assert reply.headers["Location"].endswith("/staff/sign-in")
        reply = client.post("/staff/reports", {"programme": "BSNADM01"})
        assert reply.status_code == 302
    assert reports_count(settings) == 0


# --- The bin for reports (constitution 5.14) --------------------------------
# The task of 13 October 2026 made deletion "the report disappears": its file
# and its rows. Since the report bin of 15 October 2026,
# "Can the deleted  reports go first to the bin now and then be deleted?",
# a report staff remove waits in the bin with its file and its rows, not a
# kept report, its students waiting for a report again at once ("Yes, right
# away"); it can be put back as it was, and only a report in the bin is
# deleted. Who removed, put back and deleted which report and when is kept;
# a deleted report's number is never given again.

@hsettings(max_examples=40, deadline=None)
@given(st.lists(st.one_of(
    st.tuples(st.just("generate"), st.sampled_from(PROGRAMMES)),
    st.tuples(st.sampled_from(("remove", "restore", "delete")),
              st.integers(0, 5))), min_size=2, max_size=10),
    st.lists(generated_application(), max_size=6))
def test_the_report_bin_over_any_sequence(operations, applications):
    with Data() as d:
        # One that qualifies in each programme, so every report is made.
        sure = [{"programme": programme, "grad_type": "thesis",
                 "submitted": True, "name": "Ана Тест", "fn": "F1",
                 "requests": [("confirm", False)], "checks": ["passed"]}
                for programme in PROGRAMMES]
        for index, generated in enumerate(sure + applications):
            insert(d, generated, index)
        kept, binned = {}, {}               # id: (programme, students)
        log, used = [], set()
        for kind, value in operations:
            if kind == "generate":
                try:
                    report_id = reports.generate(d.connection, d.settings,
                                                 d.lists, value, d.account)
                except reports.NothingToReport:
                    continue
                assert report_id not in used          # never a number again
                used.add(report_id)
                kept[report_id] = (value, [r["id"] for r in
                                           reports.qualifying(d.connection,
                                                              value)])
                continue
            there = sorted(set(kept) | set(binned))
            if not there:
                continue
            report_id = there[value % len(there)]
            file = reports.path(d.settings, reports.get(d.connection,
                                                        report_id))
            content = file.read_bytes()
            if kind == "remove":
                assert reports.remove(d.connection, report_id, d.account) \
                    == (report_id in kept)
                binned.update([(report_id, kept.pop(report_id))]
                              if report_id in kept else [])
            elif kind == "restore":
                assert reports.restore(d.connection, report_id, d.account) \
                    == (report_id in binned)
                kept.update([(report_id, binned.pop(report_id))]
                            if report_id in binned else [])
            else:
                deleted = reports.delete(d.connection, d.settings, report_id,
                                         d.account)
                assert deleted == (report_id in binned)
                if deleted:
                    assert not file.exists()
                    programme, listed = binned.pop(report_id)
                    log.append((report_id, programme, len(listed)))
                    continue
            # Removing and putting back never touch the file.
            assert file.read_bytes() == content
        event("{} deleted, {} in the bin, {} kept".format(
            len(log), len(binned), len(kept)))
        assert sorted(r["id"] for r in reports.kept(d.connection)) == \
            sorted(kept)
        assert sorted(r["id"] for r in reports.binned(d.connection)) == \
            sorted(binned)
        for report_id in kept:
            report = reports.get(d.connection, report_id)
            assert (report["binned_at"], report["binned_by"]) == (None, None)
        for report_id in binned:
            assert reports.get(d.connection, report_id)["binned_by"] == \
                d.account
        for report_id, (programme, listed) in list(kept.items()) + list(
                binned.items()):
            assert [r["application_id"] for r in reports.students(
                d.connection, report_id)] == listed
        files = (d.settings.reports_dir.iterdir()
                 if d.settings.reports_dir.exists() else [])
        assert sorted(p.name for p in files) == sorted(
            r["filename"] for r in d.connection.execute(
                "SELECT filename FROM reports"))
        deletions = reports.deletions(d.connection)
        assert [(r["report_id"], r["programme"], r["students"])
                for r in deletions] == list(reversed(log))
        assert all(r["removed_by"] is not None for r in deletions)
        assert d.connection.execute(
            "SELECT COUNT(*) FROM report_students WHERE report_id NOT IN "
            "(SELECT id FROM reports)").fetchone()[0] == 0
        # Waiting for a report: exactly the qualifying applications no kept
        # report lists; a report in the bin counts for nothing.
        from superassistant import applications as module
        listed = set(a for _, ids in kept.values() for a in ids)
        for row in d.connection.execute("SELECT * FROM applications"):
            standing = module.standing(d.connection, row)
            qualifies = row["id"] in [r["id"] for programme in PROGRAMMES
                                      for r in reports.qualifying(
                                          d.connection, programme)]
            assert standing.waiting_for_report == (
                qualifies and row["id"] not in listed)


def bin_row(page, report_id):
    """The bin's row of one report, as the page writes it."""
    part = page.split('id="bin-reports"')[1]
    return part.split('data-id="{}"'.format(report_id))[1].split("</tr>")[0]


def test_staff_remove_put_back_and_delete_a_report_in_two_steps(
        settings, app, clock):
    staff = make_account(settings)
    other = make_account(settings, username="other")
    client = staff_client(app, staff)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", name2="Бойкова",
                                name3="Вълчева", fn="F300001",
                                topic="Тема", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    client.post("/staff/reports", {"programme": "BSNADM01"})
    connection = db.connect(settings.db_path)
    report = reports.get(connection, 1)
    file = reports.path(settings, report)
    content = file.read_bytes()
    page = client.get("/staff/reports").get_data(as_text=True)
    assert 'action="/staff/reports/1/remove"' in page
    assert texts.REPORTS_REMOVE_HINT in page
    assert 'href="/staff/reports/1/delete"' not in page
    assert "чака доклад" not in client.get("/staff").get_data(as_text=True)

    # The first step: one press, and the report is in the bin, file and
    # rows kept; its student waits for a report again at once.
    clock.advance(hours=1)
    removed_at = db.now()
    reply = client.post("/staff/reports/1/remove", follow_redirects=True)
    page = reply.get_data(as_text=True)
    assert reply.request.path == "/staff/reports"
    assert texts.REPORTS_REMOVED in page and 'class="flash ok"' in page
    assert texts.REPORTS_NONE_KEPT in page and "/staff/reports/1." not in page
    assert file.read_bytes() == content
    assert reports_count(settings) == 1
    assert len(reports.students(connection, 1)) == 1
    assert "чака доклад" in client.get("/staff").get_data(as_text=True)
    from superassistant import applications
    assert not applications.in_a_report(connection, number)

    # The bin lists it, downloadable, with who removed it and when.
    page = client.get("/staff/bin").get_data(as_text=True)
    assert texts.BIN_NO_APPLICATIONS in page
    row = bin_row(page, 1)
    for shown in ("doklad-temi-1-2026-10-12.docx", "12.10.2026",
                  "БП „Администрация и управление“, РО", "<td>1</td>",
                  "{}, {}".format(staff["username"],
                                  db.local_datetime(removed_at)),
                  'href="/staff/reports/1.docx"',
                  'action="/staff/reports/1/restore"',
                  'href="/staff/reports/1/delete"'):
        assert shown in row, shown
    reply = client.get("/staff/reports/1.docx")
    assert reply.status_code == 200 and reply.data == content

    # Put back, as it was: the student is in a report again.
    reply = client.post("/staff/reports/1/restore", follow_redirects=True)
    assert reply.request.path == "/staff/reports"
    page = reply.get_data(as_text=True)
    assert texts.REPORTS_RESTORED in page and "/staff/reports/1.docx" in page
    restored = reports.get(connection, 1)
    assert (restored["binned_at"], restored["binned_by"]) == (None, None)
    assert dict(restored) == dict(report)                # as it was
    assert "чака доклад" not in client.get("/staff").get_data(as_text=True)
    assert texts.BIN_EMPTY in client.get("/staff/bin").get_data(as_text=True)

    # Removed again, by another member of staff; the second step names the
    # report, and the bin, not the reports, is where it is.
    clock.advance(hours=1)
    removed_at = db.now()
    staff_client(app, other).post("/staff/reports/1/remove")
    page = client.get("/staff/reports/1/delete").get_data(as_text=True)
    for shown in ("doklad-temi-1-2026-10-12.docx",
                  "БП „Администрация и управление“, РО", "12.10.2026",
                  "<dd>1</dd>", texts.REPORTS_DELETE_TITLE,
                  texts.REPORTS_DELETE_LEAD, texts.BIN_DELETE,
                  "other, {}".format(db.local_datetime(removed_at)),
                  'aria-current="page">{}'.format(texts.NAV_BIN),
                  'class="back" href="/staff/bin"',
                  'class="btn secondary" href="/staff/bin"'):
        assert shown in page, shown
    assert 'aria-current="page">{}'.format(texts.NAV_REPORTS) not in page
    # The reports page is reached only from the navigation.
    assert page.count('href="/staff/reports"') == 1
    assert reports.get(connection, 1) is not None       # nothing deleted yet

    clock.advance(hours=2)
    reply = client.post("/staff/reports/1/delete", follow_redirects=True)
    assert reply.request.path == "/staff/bin"
    page = reply.get_data(as_text=True)
    assert texts.REPORTS_DELETED in page and texts.BIN_EMPTY in page
    assert not file.exists()
    assert reports.get(connection, 1) is None
    assert reports_count(settings) == 0
    assert connection.execute("SELECT COUNT(*) FROM report_students"
                              ).fetchone()[0] == 0
    (deletion,) = reports.deletions(connection)
    assert (deletion["report_id"], deletion["programme"],
            deletion["generated_at"], deletion["students"],
            deletion["filename"], deletion["username"],
            deletion["deleted_at"], deletion["removed_by"],
            deletion["removed_at"]) == (
        1, "BSNADM01", report["generated_at"], 1, report["filename"],
        staff["username"], db.now(), "other", removed_at)
    assert [(r["report_id"], r["event"], r["by_account"]) for r in
            connection.execute("SELECT * FROM report_bin_events ORDER BY "
                               "id")] == [
        (1, "removed", staff["id"]), (1, "restored", staff["id"]),
        (1, "removed", other["id"])]
    log = page.split('id="report-deletions"')[1]
    for shown in ("doklad-temi-1-2026-10-12.docx", "12.10.2026",
                  "БП „Администрация и управление“, РО",
                  "other, {}".format(db.local_datetime(removed_at)),
                  "{}, {}".format(staff["username"],
                                  db.local_datetime(db.now()))):
        assert shown in log, shown
    assert 'id="report-deletions"' not in client.get(
        "/staff/reports").get_data(as_text=True)
    assert client.get("/staff/reports/1.docx").status_code == 404
    reply = client.get("/staff/reports/1/delete", follow_redirects=True)
    assert reply.request.path == "/staff/bin"
    assert texts.REPORTS_NOT_IN_BIN in reply.get_data(as_text=True)

    # The student waits for a report again, and the next report lists them
    # under a new number.
    assert "чака доклад" in client.get("/staff").get_data(as_text=True)
    client.post("/staff/reports", {"programme": "BSNADM01"})
    (kept,) = reports.kept(connection)
    assert kept["id"] == 2
    assert [r["application_id"] for r in reports.students(connection, 2)] \
        == [number]
    connection.close()


def test_only_a_report_in_the_bin_is_deleted(settings, app):
    """A page from before the bin, or one read before another member of
    staff put the report back, asks to delete a kept report: nothing is
    deleted, and staff are told so in the bin."""
    staff = make_account(settings)
    client = staff_client(app, staff)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    client.post("/staff/reports", {"programme": "BSNADM01"})
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, 1))
    assert reports.delete(connection, settings, 1, staff["id"]) is False
    for reply in (client.post("/staff/reports/1/delete"),
                  client.get("/staff/reports/1/delete")):
        assert reply.status_code == 302
        assert reply.headers["Location"].endswith("/staff/bin")
    page = client.get("/staff/bin").get_data(as_text=True)
    assert texts.REPORTS_NOT_IN_BIN in page
    # Read in the bin, put back by someone else, then pressed.
    client.post("/staff/reports/1/remove")
    assert "/staff/reports/1/delete" in client.get("/staff/bin").get_data(
        as_text=True)
    staff_client(app, make_account(settings, username="other")).post(
        "/staff/reports/1/restore")
    reply = client.post("/staff/reports/1/delete", follow_redirects=True)
    assert texts.REPORTS_NOT_IN_BIN in reply.get_data(as_text=True)
    assert [r["id"] for r in reports.kept(connection)] == [1]
    assert file.is_file() and reports.deletions(connection) == []
    connection.close()


def test_pressing_a_report_bin_button_twice_says_so_in_bulgarian(settings,
                                                                 app):
    staff = make_account(settings)
    client = staff_client(app, staff)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    client.post("/staff/reports", {"programme": "BSNADM01"})
    connection = db.connect(settings.db_path)

    def twice(url, said):
        client.post(url)
        reply = client.post(url, follow_redirects=True)
        assert reply.request.path == "/staff/bin"
        page = reply.get_data(as_text=True)
        assert said in page and 'class="flash error"' in page
    twice("/staff/reports/1/remove", texts.REPORTS_REMOVED_MEANWHILE)
    twice("/staff/reports/1/restore", texts.REPORTS_NOT_IN_BIN)
    client.post("/staff/reports/1/remove")
    twice("/staff/reports/1/delete", texts.REPORTS_NOT_IN_BIN)
    for url in ("/staff/reports/1/remove", "/staff/reports/1/restore",
                "/staff/reports/9/remove", "/staff/reports/9/restore",
                "/staff/reports/9/delete"):
        assert client.post(url).headers["Location"].endswith("/staff/bin")
    for url in ("/staff/reports/1/delete", "/staff/reports/9/delete"):
        reply = client.get(url, follow_redirects=True)
        assert reply.request.path == "/staff/bin"
        assert texts.REPORTS_NOT_IN_BIN in reply.get_data(as_text=True)
    # One row for each press that did something, and nothing else.
    assert [r["event"] for r in connection.execute(
        "SELECT event FROM report_bin_events ORDER BY id")] == [
        "removed", "restored", "removed"]
    assert len(reports.deletions(connection)) == 1
    connection.close()


@pytest.mark.parametrize("step, meanwhile, events", [
    ("delete", "restore", ["removed", "restored"]),
    ("remove", "remove", ["removed"]),
    ("restore", "restore", ["removed", "restored"]),
])
def test_each_step_reads_the_bin_under_its_own_write_lock(settings, app,
                                                          step, meanwhile,
                                                          events):
    """Another member of staff acts in the moment before this step takes
    the write lock. Each step reads whether the report is in the bin only
    once it holds the lock, so it sees that and changes nothing: a report
    put back just then is never deleted in one step, and nothing is logged
    twice."""
    staff, report_id = binned_report(settings)
    if step == "remove":
        connection = db.connect(settings.db_path)
        assert reports.restore(connection, report_id, staff["id"])
        connection.close()
        events = ["removed", "restored", "removed"]
    first = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(first, report_id))
    fired = []

    def other_first(sql):
        if sql == "BEGIN IMMEDIATE" and not fired:
            fired.append(sql)               # before this lock is taken
            other = db.connect(settings.db_path)
            try:
                assert getattr(reports, meanwhile)(other, report_id,
                                                   staff["id"])
            finally:
                other.close()
    first.set_trace_callback(other_first)
    arguments = ((first, settings) if step == "delete" else (first,)) + (
        report_id, staff["id"])
    assert getattr(reports, step)(*arguments) is False
    first.set_trace_callback(None)
    assert fired and not first.in_transaction
    assert file.is_file() and reports.deletions(first) == []
    assert [r["event"] for r in first.execute(
        "SELECT event FROM report_bin_events ORDER BY id")] == events
    in_bin = events[-1] == "removed"
    assert [r["id"] for r in (reports.binned(first) if in_bin
                              else reports.kept(first))] == [report_id]
    first.close()


def test_the_bin_lists_the_latest_removed_report_first(settings, app,
                                                       clock):
    staff = make_account(settings)
    client = staff_client(app, staff)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    client.post("/staff/reports", {"programme": "BSNADM01"})
    client.post("/staff/reports", {"programme": "BSNADM01"})
    client.post("/staff/reports/2/remove")
    clock.advance(hours=1)
    client.post("/staff/reports/1/remove")
    connection = db.connect(settings.db_path)
    assert [r["id"] for r in reports.binned(connection)] == [1, 2]
    connection.close()
    page = client.get("/staff/bin").get_data(as_text=True)
    table = page.split('id="bin-reports"')[1]
    assert table.index('data-id="1"') < table.index('data-id="2"')
    assert texts.BIN_NO_APPLICATIONS in page and texts.BIN_EMPTY not in page


def test_a_report_in_the_bin_keeps_its_number(settings, app):
    """Numbers are never given twice: a report in the bin keeps its own, a
    later report takes the next, and a deleted one's is not given again."""
    staff = make_account(settings)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    connection = db.connect(settings.db_path)
    loaded = lists.load(settings.lists_dir)

    def generate():
        return reports.generate(connection, settings, loaded, "BSNADM01",
                                staff["id"])
    assert generate() == 1
    assert reports.remove(connection, 1, staff["id"])
    assert generate() == 2
    assert reports.remove(connection, 2, staff["id"])
    assert reports.delete(connection, settings, 2, staff["id"])
    assert generate() == 3
    assert reports.restore(connection, 1, staff["id"])
    assert [r["id"] for r in reports.kept(connection)] == [3, 1]
    assert [r["filename"] for r in reports.kept(connection)] == [
        r["filename"] for r in connection.execute(
            "SELECT filename FROM reports ORDER BY id DESC")]
    connection.close()


def test_an_application_deleted_while_its_report_is_in_the_bin(settings,
                                                               app):
    """Constitution 5.15: the application's rows go from every report, the
    ones in the bin included; the report's file stays as it was generated,
    and the report can still be put back, removed and deleted."""
    from superassistant import applications
    staff = make_account(settings)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    connection = db.connect(settings.db_path)
    report_id = reports.generate(connection, settings, lists.load(
        settings.lists_dir), "BSNADM01", staff["id"])
    file = reports.path(settings, reports.get(connection, report_id))
    content = file.read_bytes()
    assert reports.remove(connection, report_id, staff["id"])
    assert applications.remove(connection, number, "other", "", staff["id"])
    assert applications.delete(connection, number, staff["id"])
    (binned,) = reports.binned(connection)
    assert (binned["id"], binned["students"]) == (report_id, 0)
    assert file.read_bytes() == content
    page = staff_client(app, staff).get("/staff/bin").get_data(as_text=True)
    assert "<td>0</td>" in bin_row(page, report_id)
    assert reports.restore(connection, report_id, staff["id"])
    assert [r["students"] for r in reports.kept(connection)] == [0]
    assert reports.remove(connection, report_id, staff["id"])
    assert reports.delete(connection, settings, report_id, staff["id"])
    assert [r["students"] for r in reports.deletions(connection)] == [0]
    assert not file.exists()
    connection.close()


def test_a_download_that_meets_the_deletion_is_not_found(settings, app,
                                                         monkeypatch):
    """The file goes between the look at the record and the reading: the
    download says not found, not an error page."""
    staff, report_id = binned_report(settings)
    client = staff_client(app, staff)
    path_type = type(settings.reports_dir)

    def gone(self):
        raise FileNotFoundError(2, "No such file or directory")
    monkeypatch.setattr(path_type, "read_bytes", gone)
    app.config.update(TESTING=False, PROPAGATE_EXCEPTIONS=False)
    url = "/staff/reports/{}.docx".format(report_id)
    assert client.get(url).status_code == 404
    monkeypatch.undo()
    # Nor is a folder where the file should be.
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))
    connection.close()
    file.unlink()
    file.mkdir()
    assert client.get(url).status_code == 404


def test_schema_4_with_reports_upgrades_keeping_them(tmp_path):
    """A database of the application bin (15 October 2026), with a kept
    report, its student and a report deleted then, becomes version 5 with
    the report kept and the deletion logged, no removal known for it."""
    path = tmp_path / "superassistant.sqlite3"
    connection = sqlite3.connect(str(path))
    connection.executescript(";".join(db.MIGRATIONS[:4])
                             + ";PRAGMA user_version = 4;")
    connection.executescript("""
        INSERT INTO accounts (username, username_key, email, email_key,
            algorithm, salt, iterations, dklen, hash, session_tag,
            created_at, approved_at) VALUES ('staff', 'staff',
            's@example.org', 's@example.org', 'none', '00', 1, 1, '00',
            'tag', 't', 't');
        INSERT INTO applications (email, email_key, created_at, updated_at,
            submitted_at, name1, fn) VALUES ('a@example.org',
            'a@example.org', 't', 't', 't', 'Ана', 'F1');
        INSERT INTO reports (id, programme, faculty, generated_at,
            generated_by, filename) VALUES (2, 'BSNADM01', 'ФДЕНО',
            '2026-10-12T09:00:00+00:00', 1, 'doklad-temi-x.docx');
        INSERT INTO report_students (report_id, position, application_id,
            fn, name, topic, supervisor) VALUES (2, 1, 1, 'F1', 'Ана',
            'Тема', 'проф. Иванов');
        INSERT INTO report_deletions (report_id, programme, generated_at,
            students, filename, deleted_by, deleted_at) VALUES (1,
            'BSNADM01', '2026-10-11T09:00:00+00:00', 3, 'doklad-temi-y.docx',
            1, '2026-10-11T10:00:00+00:00');
    """)
    connection.commit()
    connection.close()
    db.migrate(path)
    connection = db.connect(path)
    assert connection.execute("PRAGMA user_version").fetchone()[0] == len(
        db.MIGRATIONS) == 5
    (kept,) = reports.kept(connection)
    assert (kept["id"], kept["students"], kept["binned_at"],
            kept["binned_by"]) == (2, 1, None, None)
    assert reports.binned(connection) == []
    (deletion,) = reports.deletions(connection)
    assert (deletion["report_id"], deletion["students"],
            deletion["username"], deletion["removed_by"],
            deletion["removed_at"]) == (1, 3, "staff", None, None)
    from superassistant import applications
    assert applications.in_a_report(connection, 1)
    connection.close()


def test_a_deletion_from_before_the_bin_shows_no_removal(settings, app):
    staff = make_account(settings)
    connection = db.connect(settings.db_path)
    connection.execute(
        "INSERT INTO report_deletions (report_id, programme, generated_at, "
        "students, filename, deleted_by, deleted_at) VALUES (1, "
        "'BSNADM01', '2026-10-11T09:00:00+00:00', 3, 'doklad-temi-y.docx', "
        "?, '2026-10-11T10:00:00+00:00')", (staff["id"],))
    connection.close()
    page = staff_client(app, staff).get("/staff/bin").get_data(as_text=True)
    cells = page.split('id="report-deletions"')[1].split("<tbody>")[1].split(
        "</tr>")[0].split("<td>")
    assert cells[1].startswith("11.10.2026") and cells[3].startswith("3<")
    assert cells[4].startswith("doklad-temi-1-2026-10-11.docx")
    assert cells[5].startswith("</td>")                    # no removal
    assert cells[6].startswith("{}, ".format(staff["username"]))


def test_a_missing_file_does_not_stop_the_deletion(settings, app):
    staff = make_account(settings)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    connection = db.connect(settings.db_path)
    loaded = lists.load(settings.lists_dir)
    report_id = reports.generate(connection, settings, loaded, "BSNADM01",
                                 staff["id"])
    assert reports.remove(connection, report_id, staff["id"])
    reports.path(settings, reports.get(connection, report_id)).unlink()
    assert reports.delete(connection, settings, report_id, staff["id"])
    assert reports.binned(connection) == [] and reports.kept(connection) == []
    assert len(reports.deletions(connection)) == 1
    assert reports.delete(connection, settings, report_id, staff["id"]) \
        is False
    connection.close()


def binned_report(settings):
    """A staff account and one report in the bin listing one student: the
    deletion's second step can take it."""
    staff = make_account(settings)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    connection = db.connect(settings.db_path)
    try:
        report_id = reports.generate(connection, settings, lists.load(
            settings.lists_dir), "BSNADM01", staff["id"])
        assert reports.remove(connection, report_id, staff["id"])
    finally:
        connection.close()
    return staff, report_id


def test_a_commit_held_off_by_a_reader_keeps_the_report_and_its_file(
        settings, app):
    """Codex T13-C1: the file went before the COMMIT, which another
    connection still reading can hold off past the busy timeout; the rows
    were then rolled back, leaving a kept report without its file. A real
    reader lock, with a short busy timeout in place of the 30 seconds."""
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    connection.execute("PRAGMA busy_timeout = 200")
    file = reports.path(settings, reports.get(connection, report_id))
    content = file.read_bytes()
    reader = db.connect(settings.db_path)
    reader.execute("BEGIN")
    assert reader.execute("SELECT COUNT(*) FROM reports").fetchone()[0] == 1
    with pytest.raises(sqlite3.OperationalError) as raised:
        reports.delete(connection, settings, report_id, staff["id"])
    assert "locked" in str(raised.value)
    assert file.is_file() and file.read_bytes() == content
    assert not aside(file).exists()
    assert not connection.in_transaction            # rolled back at once
    reader.execute("ROLLBACK")
    reader.close()
    other = db.connect(settings.db_path)
    for seen in (connection, other):
        assert [r["id"] for r in reports.binned(seen)] == [report_id]
        assert len(reports.students(seen, report_id)) == 1
        assert reports.deletions(seen) == []
    other.close()
    # The report, still in the bin, still downloads.
    reply = staff_client(app, staff).get(
        "/staff/reports/{}.docx".format(report_id))
    assert reply.status_code == 200 and reply.data == content
    # Once nobody is reading, the same connection deletes it.
    assert reports.delete(connection, settings, report_id, staff["id"])
    assert not file.exists() and reports.binned(connection) == []
    assert len(reports.deletions(connection)) == 1
    connection.close()


def test_a_failure_inside_the_deletion_keeps_everything_and_names_itself(
        settings, app):
    """A write that fails inside the transaction is rolled back by it; the
    caller sees that error, not a second ROLLBACK's, and nothing changes."""
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))
    with pytest.raises(sqlite3.IntegrityError):         # no such account
        reports.delete(connection, settings, report_id, 999999)
    assert not connection.in_transaction
    assert file.is_file() and not aside(file).exists()
    assert [r["id"] for r in reports.binned(connection)] == [report_id]
    assert len(reports.students(connection, report_id)) == 1
    assert reports.deletions(connection) == []
    connection.close()


def aside(file):
    return file.with_name(file.name + reports.ASIDE)


def counts(settings):
    connection = db.connect(settings.db_path)
    try:
        return tuple(connection.execute("SELECT COUNT(*) FROM " + table
                                        ).fetchone()[0]
                     for table in ("reports", "report_students",
                                   "report_deletions"))
    finally:
        connection.close()


def test_a_transaction_sqlite_rolls_back_itself_reports_its_own_error(
        settings, app):
    """An interrupt inside the deletion makes SQLite roll the transaction
    back by itself; the caller sees that error, not a second ROLLBACK's
    ("cannot rollback - no transaction is active"), and nothing changes."""
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))

    def interrupt(sql):
        if sql.startswith("DELETE FROM reports"):
            connection.interrupt()
    connection.set_trace_callback(interrupt)
    with pytest.raises(sqlite3.OperationalError) as raised:
        reports.delete(connection, settings, report_id, staff["id"])
    connection.set_trace_callback(None)
    assert "interrupt" in str(raised.value)
    assert not connection.in_transaction
    assert file.is_file() and not aside(file).exists()
    assert counts(settings) == (1, 1, 0)
    connection.close()


class CommitAs(object):
    """A connection whose COMMIT is replaced, standing for what a real one
    does rarely: SQLite rolling back by itself on an I/O error, or an error
    arriving after the COMMIT has succeeded."""

    def __init__(self, real, commit):
        self.real, self.commit = real, commit

    def execute(self, sql, *args):
        if sql == "COMMIT":
            return self.commit(self.real)
        return self.real.execute(sql, *args)

    @property
    def in_transaction(self):
        return self.real.in_transaction


def rolled_back_by_sqlite(real):
    real.execute("ROLLBACK")
    raise sqlite3.OperationalError("disk I/O error")


def committed_then_interrupted(real):
    real.execute("COMMIT")
    raise KeyboardInterrupt()


@pytest.mark.parametrize("commit, error, kept", [
    (rolled_back_by_sqlite, sqlite3.OperationalError, True),
    (committed_then_interrupted, KeyboardInterrupt, False),
])
def test_the_records_decide_whether_the_file_goes_back(settings, app, commit,
                                                       error, kept):
    """The records, read again under the write lock, decide: after
    SQLite's own rollback the report is kept and its file goes back; after
    a COMMIT that succeeded it is gone, and so is its file, never left
    under its name, where tidy would not see it."""
    staff, report_id = binned_report(settings)
    real = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(real, report_id))
    content = file.read_bytes()
    with pytest.raises(error):
        reports.delete(CommitAs(real, commit), settings, report_id,
                       staff["id"])
    assert not real.in_transaction
    if kept:
        assert counts(settings) == (1, 1, 0)
        assert file.read_bytes() == content and not aside(file).exists()
    else:
        assert counts(settings) == (0, 0, 1)
        assert not file.exists() and not aside(file).exists()
    real.close()


def test_a_second_deletion_cannot_slip_between_the_look_and_the_put_back(
        settings, app, monkeypatch):
    """Codex T13-C3. The first deletion's transaction ends without its
    ROLLBACK (SQLite rolling back by itself at the COMMIT), so it reads the
    records again to settle its file. A second deletion of the same report,
    on a second connection, is started at the moment the first would put
    the file back, and held after its own COMMIT until the first has
    settled. Unlocked, the second commits in between and the file of a
    deleted report goes back under its name; under the write lock the
    second waits, and the report ends deleted with its file gone."""
    import threading
    staff, report_id = binned_report(settings)
    first = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(first, report_id))
    second_committed = threading.Event()
    first_settled = threading.Event()
    outcome = {}

    def run_second():
        second = db.connect(settings.db_path)   # its own thread's connection
        try:
            outcome["second"] = reports.delete(second, settings, report_id,
                                               staff["id"])
        except BaseException as exc:            # reported by the assertion
            outcome["second"] = exc
        outcome["open"] = second.in_transaction
        second.close()
    worker = threading.Thread(target=run_second)
    path_type = type(file)
    real_replace, real_unlink = path_type.replace, path_type.unlink

    def replace(self, target):
        putting_back = (threading.current_thread() is not worker
                        and self.name.endswith(reports.ASIDE))
        if putting_back:
            worker.start()
            second_committed.wait(timeout=1)
        try:
            return real_replace(self, target)
        finally:
            if putting_back:
                first_settled.set()

    def unlink(self, *args, **kwargs):
        if (threading.current_thread() is worker
                and self.name.endswith(reports.ASIDE)):
            second_committed.set()
            first_settled.wait(timeout=5)
        return real_unlink(self, *args, **kwargs)
    monkeypatch.setattr(path_type, "replace", replace)
    monkeypatch.setattr(path_type, "unlink", unlink)
    with pytest.raises(sqlite3.OperationalError) as raised:
        reports.delete(CommitAs(first, rolled_back_by_sqlite), settings,
                       report_id, staff["id"])
    worker.join(timeout=60)
    monkeypatch.undo()
    assert not worker.is_alive()
    assert "disk I/O" in str(raised.value)
    assert outcome["second"] is True
    assert not first.in_transaction and outcome["open"] is False
    assert counts(settings) == (0, 0, 1)
    assert not file.exists() and not aside(file).exists()
    first.close()


def test_a_file_that_cannot_be_removed_keeps_the_report(settings, app,
                                                        monkeypatch):
    """Codex T13-C2: the deletion is whole or nothing. A file that cannot
    be moved aside, as removing it needs the same rights, keeps the report,
    its rows and its file, logs nothing, and the deletion can be retried."""
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))

    def refuse(self, *args, **kwargs):
        raise PermissionError(13, "Permission denied")
    monkeypatch.setattr(type(settings.reports_dir), "replace", refuse)
    monkeypatch.setattr(type(settings.reports_dir), "unlink", refuse)
    with pytest.raises(PermissionError):
        reports.delete(connection, settings, report_id, staff["id"])
    monkeypatch.undo()
    assert not connection.in_transaction
    assert file.is_file() and not aside(file).exists()
    assert counts(settings) == (1, 1, 0)
    assert reports.delete(connection, settings, report_id, staff["id"])
    assert not file.exists() and not aside(file).exists()
    assert counts(settings) == (0, 0, 1)
    connection.close()


@pytest.mark.skipif(os.geteuid() == 0, reason="root ignores folder rights")
def test_a_reports_folder_that_refuses_changes_keeps_the_report(settings,
                                                                app):
    """Codex T13-C2 with real folder rights, through the page as the host
    serves it: an error page, nothing changed, the report still downloads,
    and once the rights are back the deletion goes through."""
    app.config.update(TESTING=False, PROPAGATE_EXCEPTIONS=False)
    staff, report_id = binned_report(settings)
    client = staff_client(app, staff)
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))
    connection.close()
    content = file.read_bytes()
    url = "/staff/reports/{}".format(report_id)
    mode = settings.reports_dir.stat().st_mode & 0o777
    settings.reports_dir.chmod(0o500)
    try:
        assert client.post(url + "/delete").status_code == 500
    finally:
        settings.reports_dir.chmod(mode)
    assert counts(settings) == (1, 1, 0)
    assert file.read_bytes() == content and not aside(file).exists()
    reply = client.get(url + ".docx")
    assert reply.status_code == 200 and reply.data == content
    assert client.post(url + "/delete").status_code == 302
    assert counts(settings) == (0, 0, 1)
    assert not file.exists() and not aside(file).exists()
    assert client.get(url + ".docx").status_code == 404


def test_the_file_is_aside_at_the_commit_and_back_before_the_rollback(
        settings, app):
    """The order that makes a failed deletion whole: the file is moved
    before the COMMIT, and put back before the ROLLBACK gives up the write
    lock, which tidy waits for."""
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    connection.execute("PRAGMA busy_timeout = 200")
    file = reports.path(settings, reports.get(connection, report_id))
    seen = []
    connection.set_trace_callback(lambda sql: seen.append(
        (sql, file.exists(), aside(file).exists())))
    reader = db.connect(settings.db_path)
    reader.execute("BEGIN")
    reader.execute("SELECT COUNT(*) FROM reports").fetchone()
    with pytest.raises(sqlite3.OperationalError):
        reports.delete(connection, settings, report_id, staff["id"])
    connection.set_trace_callback(None)
    reader.execute("ROLLBACK")
    reader.close()
    assert ("COMMIT", False, True) in seen
    assert seen[-1] == ("ROLLBACK", True, False)
    # The transaction still open, the file goes back under the lock it
    # already holds, without letting it go and taking it again.
    assert [sql for sql, _, _ in seen if sql in (
        "BEGIN IMMEDIATE", "COMMIT", "ROLLBACK")] == [
        "BEGIN IMMEDIATE", "COMMIT", "ROLLBACK"]
    assert counts(settings) == (1, 1, 0)
    connection.close()


def test_a_start_finishes_what_a_stopped_deletion_left(settings, app,
                                                       monkeypatch):
    """A process stopped between setting the file aside and the COMMIT
    leaves the file of a report still there aside; one stopped after the
    COMMIT leaves
    the file of a deleted report. The next start puts back the first and
    removes the second."""
    from superassistant.web import create_app
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    kept_file = reports.path(settings, reports.get(connection, report_id))
    content = kept_file.read_bytes()
    kept_file.replace(aside(kept_file))               # stopped before COMMIT
    # A kept report's file aside too: only the file of a report whose row
    # is gone is removed, the report in the bin or not.
    other_id = reports.generate(connection, settings, lists.load(
        settings.lists_dir), "BSNADM01", staff["id"])
    other_file = reports.path(settings, reports.get(connection, other_id))
    other_content = other_file.read_bytes()
    other_file.replace(aside(other_file))
    deleted_id = reports.generate(connection, settings, lists.load(
        settings.lists_dir), "BSNADM01", staff["id"])
    deleted_file = reports.path(settings, reports.get(connection,
                                                      deleted_id))
    assert reports.remove(connection, deleted_id, staff["id"])
    def refuse(self, *args, **kwargs):
        raise PermissionError(13, "Permission denied")
    # The removal after the COMMIT fails (or the process stops before it):
    # the deletion stands and reports success.
    monkeypatch.setattr(type(settings.reports_dir), "unlink", refuse)
    assert reports.delete(connection, settings, deleted_id, staff["id"])
    monkeypatch.undo()
    connection.close()
    assert aside(deleted_file).exists() and not deleted_file.exists()
    assert counts(settings) == (2, 2, 1)
    folder = settings.reports_dir / ("folder" + reports.ASIDE)
    folder.mkdir()                                  # not a file: left alone
    started = create_app(settings)
    assert kept_file.read_bytes() == content and not aside(kept_file).exists()
    assert other_file.read_bytes() == other_content
    assert not aside(other_file).exists()
    assert not aside(deleted_file).exists() and not deleted_file.exists()
    assert sorted(p.name for p in settings.reports_dir.iterdir()) == sorted(
        [kept_file.name, other_file.name, folder.name])
    reply = staff_client(started, staff).get(
        "/staff/reports/{}.docx".format(report_id))
    assert reply.status_code == 200 and reply.data == content


def test_tidying_waits_for_a_deletion_in_another_process(settings, app):
    """tidy holds the write lock while it looks, so a deletion still running
    elsewhere, which holds it until its COMMIT or ROLLBACK, is waited for."""
    other = db.connect(settings.db_path)
    other.execute("BEGIN IMMEDIATE")
    connection = db.connect(settings.db_path)
    connection.execute("PRAGMA busy_timeout = 200")
    with pytest.raises(sqlite3.OperationalError) as raised:
        reports.tidy(connection, settings)
    assert "locked" in str(raised.value)
    other.execute("ROLLBACK")
    other.close()
    reports.tidy(connection, settings)
    assert not connection.in_transaction
    connection.close()


@pytest.mark.skipif(os.geteuid() == 0, reason="root ignores folder rights")
def test_a_reports_folder_that_cannot_be_tidied_stops_the_start(settings,
                                                                app):
    from superassistant.web import StartupError, create_app
    staff, report_id = binned_report(settings)
    connection = db.connect(settings.db_path)
    file = reports.path(settings, reports.get(connection, report_id))
    connection.close()
    file.replace(aside(file))
    mode = settings.reports_dir.stat().st_mode & 0o777
    settings.reports_dir.chmod(0o500)
    try:
        with pytest.raises(StartupError) as raised:
            create_app(settings)
    finally:
        settings.reports_dir.chmod(mode)
    assert "reports folder" in str(raised.value)
    assert aside(file).exists()
    create_app(settings)
    assert file.is_file() and not aside(file).exists()


def test_only_approved_staff_may_delete_a_report(settings, app):
    staff = make_account(settings)
    number = insert_application(settings, "a@example.org",
                                grad_type="thesis", programme="BSNADM01",
                                name1="Ана", fn="F1", professor="ivanov")
    insert_request(settings, number, answer="confirm")
    insert_check(settings, number, "passed", staff["id"])
    connection = db.connect(settings.db_path)
    report_id = reports.generate(connection, settings, lists.load(
        settings.lists_dir), "BSNADM01", staff["id"])
    file = reports.path(settings, reports.get(connection, report_id))
    waiting = make_account(settings, username="waiting", approved=False)
    student = Client(app.test_client())
    with student.raw.session_transaction() as session:
        session["student"] = "a@example.org"
    base = "/staff/reports/{}/".format(report_id)
    for client in (Client(app.test_client()), student,
                   staff_client(app, waiting)):
        for reply in (client.get(base + "delete"),
                      client.post(base + "delete"),
                      client.post(base + "remove"),
                      client.post(base + "restore")):
            assert reply.status_code == 302
            assert reply.headers["Location"].endswith("/staff/sign-in")
            assert "Ана" not in reply.get_data(as_text=True)
    assert [r["id"] for r in reports.kept(connection)] == [report_id]
    # In the bin, the same: nobody but approved staff puts it back or
    # deletes it.
    assert reports.remove(connection, report_id, staff["id"])
    for client in (Client(app.test_client()), student,
                   staff_client(app, waiting)):
        for reply in (client.get(base + "delete"),
                      client.post(base + "delete"),
                      client.post(base + "restore")):
            assert reply.status_code == 302
            assert reply.headers["Location"].endswith("/staff/sign-in")
    assert [r["id"] for r in reports.binned(connection)] == [report_id]
    assert file.is_file() and reports.deletions(connection) == []
    assert [r["event"] for r in connection.execute(
        "SELECT event FROM report_bin_events")] == ["removed"]
    connection.close()
