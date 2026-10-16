"""The Word files, read back from the .docx (NOW.md requirements 4, 6 and 7;
proof 2, by content): the заявление-декларация for each graduation type,
Приложение 2 and the topic-approval report, over generated names, topics,
proposals and lists, in any characters, mixed Latin and Cyrillic included.

Generated text leaves out what a Word file cannot hold (control characters
other than a line break in a proposal): the application refuses those at
entry. Everything is invented.
"""

import io
import re
import string
import zipfile
from datetime import datetime, timezone
from zoneinfo import ZoneInfo

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from hypothesis import given, settings as hsettings, strategies as st

from superassistant import documents, lists, texts
from superassistant.web import create_app
from tests.conftest import (Client, insert_application, make_settings,
                            prepare_data_dir)

TEMPLATES = (documents.ZAIAVLENIE, documents.PRILOZHENIE2,
             documents.DOKLAD_TEMI)
CYRILLIC = ("абвгдежзийклмнопрстуфхцчшщъьюяАБВГДЕЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЬЮЯ")
SPECIAL = " -.,:;/()„“\"'&<>{}%_#№"
ANY_CHARACTER = st.characters(blacklist_categories=("Cc", "Cs"),
                              blacklist_characters="￾￿")


def line(min_size=1, max_size=40):
    """One line as the application stores it: no space at either end."""
    return st.text(alphabet=st.one_of(
        st.sampled_from(CYRILLIC + string.ascii_letters + string.digits
                        + SPECIAL), ANY_CHARACTER),
        min_size=min_size, max_size=max_size).map(str.strip).filter(
            lambda s: len(s) >= min_size)


def proposal():
    """A proposal section: lines joined by line breaks, empty lines too."""
    return st.lists(line(min_size=0, max_size=30), min_size=1,
                    max_size=5).map("\n".join).map(str.strip)


MOMENTS = st.datetimes(min_value=datetime(2020, 1, 1),
                       max_value=datetime(2035, 12, 31),
                       timezones=st.just(timezone.utc))


def stamp(moment):
    return moment.isoformat(timespec="seconds")


def sofia_date(moment):
    return moment.astimezone(ZoneInfo("Europe/Sofia")).strftime("%d.%m.%Y")


@st.composite
def generated_lists(draw):
    """A Lists object: titles, one faculty per degree and professors."""
    titles = draw(st.lists(line(max_size=12), min_size=1, max_size=4,
                           unique=True))
    faculties = {}
    for code in ("F1", "F2"):
        faculties[code] = {"name": draw(line()),
                           "dean_title": draw(st.sampled_from(titles)),
                           "dean": draw(line())}
    programmes = {}
    for code, degree in (("P-B", "БП"), ("P-M", "МП")):
        programmes[code] = {"degree": degree, "name": draw(line()),
                            "form": draw(st.sampled_from(lists.FORMS)),
                            "faculty": draw(st.sampled_from(["F1", "F2"]))}
    professors = {"prof": {"title": draw(st.sampled_from(titles)),
                           "name": draw(line()),
                           "email": "prof@example.org"}}
    return lists.Lists(titles, faculties, programmes, professors)


@st.composite
def application(draw, the_lists, grad_type):
    listed = draw(st.booleans())
    record = {
        "id": draw(st.integers(1, 10 ** 6)),
        "email": draw(line()), "grad_type": grad_type,
        "name1": draw(line()), "name2": draw(line()), "name3": draw(line()),
        "fn": draw(line(max_size=12)), "phone": draw(line(max_size=15)),
        "programme": draw(st.sampled_from(sorted(the_lists.programmes))),
        "topic": draw(line(max_size=120)),
        "average_grade": draw(line(max_size=6)),
        "professor": "prof" if listed else "",
        "sup_title": "" if listed else draw(st.sampled_from(
            the_lists.titles)),
        "sup_name": "" if listed else draw(line()),
        "sup_email": "", "submitted_at": stamp(draw(MOMENTS)),
    }
    for field in ("rationale", "aims", "object_subject", "methodology",
                  "information_base", "plan", "schedule"):
        record[field] = draw(proposal())
    return record


def read(data):
    return docx.Document(io.BytesIO(data))


def every_paragraph(document):
    """Body paragraphs, then the paragraphs of each table cell, in order."""
    out = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                out.extend(cell.paragraphs)
    return out


def template_paragraphs(name):
    return every_paragraph(docx.Document(str(documents.TEMPLATES / name)))


def slots(name):
    """{tag: index} of the paragraph each {{ tag }} sits in, and the indexes
    of the paragraphs that hold no tag at all (and so must come out as the
    template has them)."""
    tags, plain = {}, []
    for index, paragraph in enumerate(template_paragraphs(name)):
        found = re.findall(r"{{ (\w+) }}", paragraph.text)
        for tag in found:
            tags[tag] = index
        if not found:
            plain.append(index)
    return tags, plain


def assert_untouched(name, document):
    """Every paragraph without a tag reads exactly as in the template."""
    _, plain = slots(name)
    before = template_paragraphs(name)
    after = every_paragraph(document)
    for index in plain:
        assert after[index].text == before[index].text


def text_nodes(paragraph):
    return "".join(t.text or "" for t in paragraph._p.iter(qn("w:t")))


def names(record):
    return " ".join([record["name1"], record["name2"], record["name3"]])


def supervisor(record, the_lists):
    if record["professor"]:
        professor = the_lists.professors[record["professor"]]
        return professor["title"] + " " + professor["name"]
    return record["sup_title"] + " " + record["sup_name"]


SIGNATURE = "         Подпис: …………………………………."


# --- The заявление-декларация ----------------------------------------------

@hsettings(max_examples=60, deadline=None)
@given(st.data())
def test_thesis_zaiavlenie_holds_what_the_records_hold(data):
    the_lists = data.draw(generated_lists())
    record = data.draw(application(the_lists, "thesis"))
    document = read(documents.zaiavlenie(record, the_lists))
    tags, _ = slots(documents.ZAIAVLENIE)
    paragraphs = every_paragraph(document)

    def at(tag):
        return paragraphs[tags[tag]].text

    programme = the_lists.programmes[record["programme"]]
    adjective = ("бакалавърска" if programme["degree"] == "БП"
                 else "магистърска")
    submitted = sofia_date(datetime.fromisoformat(record["submitted_at"]))
    assert at("names") == "Трите имена: " + names(record)
    assert at("fn") == (record["fn"] + " тел.: " + record["phone"]
                        + " e-mail: " + record["email"])
    assert at("programme") == "Програма: {} „{}“, {}".format(
        programme["degree"], programme["name"], programme["form"])
    assert at("grad_type") == ("Моля да бъда допуснат(а) до дипломиране "
                               "чрез " + adjective + " теза")
    assert at("topic") == record["topic"]
    # No empty lines under a filled topic: the supervisor's line follows it.
    assert tags["supervisor"] == tags["topic"] + 1
    assert at("supervisor") == ("С научен (художествен) ръководител: "
                                + supervisor(record, the_lists))
    assert at("average_grade") == ("среден успех от следването: "
                                   + record["average_grade"])
    assert at("thesis_date") == ("Дата на подаване на заявлението: "
                                 + submitted + SIGNATURE)
    # Section 1, the state exam's, stays to be filled by hand.
    assert at("exam_date") == ("Дата на подаване на заявлението: "
                               + documents.DOTS["date"] + SIGNATURE)
    assert_untouched(documents.ZAIAVLENIE, document)


@hsettings(max_examples=60, deadline=None)
@given(st.data())
def test_state_exam_zaiavlenie_fills_names_fn_and_date_only(data):
    the_lists = data.draw(generated_lists())
    # Thesis fields left over in the record must not reach the form.
    record = data.draw(application(the_lists, "exam"))
    document = read(documents.zaiavlenie(record, the_lists))
    tags, _ = slots(documents.ZAIAVLENIE)
    paragraphs = every_paragraph(document)

    def at(tag):
        return paragraphs[tags[tag]].text

    dots = documents.DOTS
    submitted = sofia_date(datetime.fromisoformat(record["submitted_at"]))
    assert at("names") == "Трите имена: " + names(record)
    assert at("fn") == (record["fn"] + " тел.: " + dots["phone"]
                        + " e-mail: " + record["email"])
    assert at("programme") == "Програма: " + dots["programme"]
    assert at("grad_type") == ("Моля да бъда допуснат(а) до дипломиране "
                               "чрез държавен изпит")
    assert at("exam_date") == ("Дата на подаване на заявлението: "
                               + submitted + SIGNATURE)
    # A blank topic keeps the sample's three dotted lines, in one paragraph.
    assert at("topic") == "\n".join([dots["topic"]] * 3)
    assert len(paragraphs[tags["topic"]]._p.findall(".//" + qn("w:br"))) == 2
    assert at("supervisor") == ("С научен (художествен) ръководител: "
                                + dots["supervisor"])
    assert at("average_grade") == ("среден успех от следването: "
                                   + dots["average_grade"])
    assert at("thesis_date") == ("Дата на подаване на заявлението: "
                                 + dots["date"] + SIGNATURE)
    assert_untouched(documents.ZAIAVLENIE, document)


def test_zaiavlenie_thesis_box_goes_with_its_heading():
    """The thesis box may not split across pages (the sample's own layout):
    when it does not fit on page 1, its heading goes with it, rather than
    staying alone at the foot of the page, and the topic is one paragraph,
    with no empty lines after it (seen in Word on 16 October 2026)."""
    template = docx.Document(str(documents.TEMPLATES / documents.ZAIAVLENIE))
    body = list(template.element.body.iterchildren())
    (heading,) = [el for el in body if el.tag == qn("w:p") and "".join(
        t.text or "" for t in el.iter(qn("w:t"))).startswith(
            "2.Попълва се при завършване с бакалавърска")]
    assert heading.find(qn("w:pPr")).find(qn("w:keepNext")) is not None
    box = body[body.index(heading) + 1]
    assert box.tag == qn("w:tbl")
    (row,) = box.findall(qn("w:tr"))
    assert row.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None
    cell = [p.text for p in template.tables[1].cell(0, 0).paragraphs]
    topic = cell.index("{{ topic }}")
    assert cell[topic + 1].startswith("С научен (художествен) ръководител:")


# --- Приложение 2 ----------------------------------------------------------

@hsettings(max_examples=60, deadline=None)
@given(st.data())
def test_prilozhenie2_keeps_line_breaks_and_leaves_dates_blank(data):
    the_lists = data.draw(generated_lists())
    record = data.draw(application(the_lists, "thesis"))
    document = read(documents.prilozhenie2(record, the_lists))
    tags, _ = slots(documents.PRILOZHENIE2)
    paragraphs = every_paragraph(document)

    degree = the_lists.programmes[record["programme"]]["degree"]
    heading = "БАКАЛАВЪРСКА" if degree == "БП" else "МАГИСТЪРСКА"
    assert paragraphs[tags["thesis_heading"]].text == heading + " ТЕЗА"
    assert paragraphs[tags["names"]].text == ("На студента " + names(record)
                                              + " " + record["fn"])
    assert paragraphs[tags["topic"]].text == "ТЕМА " + record["topic"]
    for field in ("rationale", "aims", "object_subject", "methodology",
                  "information_base", "plan", "schedule"):
        paragraph = paragraphs[tags[field]]
        expected = record[field] or documents.DOTS["section"]
        assert paragraph.text == expected
        # Each line break is a line break in Word, not a character in a run.
        assert "\n" not in text_nodes(paragraph)
        assert (len(paragraph._p.findall(".//" + qn("w:br")))
                == expected.count("\n"))
    assert_untouched(documents.PRILOZHENIE2, document)
    handed_in = [p.text for p in paragraphs
                 if p.text.startswith("Дата на предаване:")]
    assert len(handed_in) == 1
    assert not re.search(r"\d", handed_in[0].split("\t")[0])


def test_prilozhenie2_does_not_stretch_a_line_ending_in_a_line_break():
    """The sections are justified, as in the sample; Word's option "Don't
    expand character spaces on a line ending with SHIFT+RETURN" keeps a line
    the student ended with Enter at its own width („Глава 1“ was printed as
    „Глава … 1“ across the page, 16 October 2026)."""
    template = docx.Document(str(documents.TEMPLATES
                                 / documents.PRILOZHENIE2))
    compat = template.settings.element.find(qn("w:compat"))
    assert compat is not None
    assert compat.find(qn("w:doNotExpandShiftReturn")) is not None
    tags, _ = slots(documents.PRILOZHENIE2)
    paragraphs = template_paragraphs(documents.PRILOZHENIE2)
    for field in ("rationale", "aims", "object_subject", "methodology",
                  "information_base", "plan", "schedule"):
        assert paragraphs[tags[field]].alignment == WD_ALIGN_PARAGRAPH.JUSTIFY


# Advance widths in Times New Roman, in thousandths of an em, regular and
# bold alike: what decides where the date line's tabs land.
TIMES_WIDTHS = {" ": 250, "…": 1000, ".": 250, "г": 500,
                **{digit: 500 for digit in "0123456789"}}
SAMPLE_DATE_LINE = ["Дата на предаване: ", "…………….", " 2025", " г."]


def test_prilozhenie2_date_line_keeps_the_samples_width():
    """The supervisor's slot is placed with tabs after the date, so the text
    before the first tab must be as wide as in the department's sample, or
    Word drops the slot to the next line's left edge (seen in Word on 10
    October 2026, when longer dots had replaced "……………. 2025"). Only the
    year is replaced, by dots of the same width, and no digit is left."""
    template = docx.Document(str(documents.TEMPLATES
                                 / documents.PRILOZHENIE2))
    (line,) = [p for p in template.paragraphs
               if p.text.startswith("Дата на предаване:")]
    before_tab = []
    for run in line.runs:
        if "\t" in run.text:
            break
        before_tab.append(run.text)
    assert before_tab[0] == SAMPLE_DATE_LINE[0]
    assert before_tab[1] == SAMPLE_DATE_LINE[1]

    def width(text):
        return sum(TIMES_WIDTHS[ch] for ch in text)

    assert [width(t) for t in before_tab[1:]] == \
        [width(t) for t in SAMPLE_DATE_LINE[1:]]
    assert not re.search(r"\d", "".join(before_tab))


def test_prilozhenie2_for_an_unlisted_programme_names_both_degrees():
    the_lists = lists.Lists(["доц. д-р"], {}, {}, {})
    record = {"name1": "Ана", "name2": "Бойкова", "name3": "Вълчева",
              "fn": "F1", "programme": "", "topic": "Тема",
              "rationale": "", "aims": "", "object_subject": "",
              "methodology": "", "information_base": "", "plan": "",
              "schedule": ""}
    document = read(documents.prilozhenie2(record, the_lists))
    tags, _ = slots(documents.PRILOZHENIE2)
    assert (every_paragraph(document)[tags["thesis_heading"]].text
            == texts.DOC_THESIS_EITHER + " ТЕЗА")


# --- The topic-approval report ---------------------------------------------

REPORT_HEADER = ["№", "F№", "Име", "Тема", "Научен ръководител", "Рецензент"]


@st.composite
def report_rows(draw):
    count = draw(st.integers(1, 5))
    return [{"position": position, "application_id": position,
             "fn": draw(line(max_size=12)), "name": draw(line(max_size=60)),
             "topic": draw(line(max_size=150)),
             "supervisor": draw(line(max_size=40))}
            for position in range(1, count + 1)]


@hsettings(max_examples=60, deadline=None)
@given(st.data())
def test_report_addressee_wording_date_and_cells(data):
    the_lists = data.draw(generated_lists())
    code = data.draw(st.sampled_from(sorted(the_lists.programmes)))
    rows = data.draw(report_rows())
    moment = data.draw(MOMENTS)
    document = read(documents.doklad_temi(code, rows, stamp(moment),
                                          the_lists))
    programme = the_lists.programmes[code]
    faculty = the_lists.faculties[programme["faculty"]]
    adjective = ("бакалавърска" if programme["degree"] == "БП"
                 else "магистърска")
    body = [p.text for p in document.paragraphs]

    assert body[0] == "ДО"
    assert body[1] == (faculty["dean_title"] + " " + faculty["dean"]).upper()
    assert body[2] == "ДЕКАН НА " + faculty["name"].upper()
    assert body[3] == "Изх. № ……. /……...… г."
    assert [b.strip() for b in body[4:7]] == ["ДОКЛАД", "за",
                                              "утвърждаване на теми"]
    intro = [b for b in body if b.startswith("Студенти от ")]
    assert intro == ["Студенти от {} „{}“, изпълнили условията за "
                     "допускане до дипломиране с {} теза:".format(
                         programme["degree"], programme["name"], adjective)]
    dated = [b for b in body if b.startswith("Дата:")]
    assert len(dated) == 1
    assert dated[0].rstrip("\t") == "Дата: " + sofia_date(moment) + " г."

    (table,) = document.tables
    cells = [[cell.text for cell in row.cells] for row in table.rows]
    assert cells[0] == REPORT_HEADER
    assert cells[1:] == [[str(row["position"]), row["fn"], row["name"],
                          row["topic"], row["supervisor"], ""]
                         for row in rows]
    assert_untouched_report(document)


def assert_untouched_report(document):
    """Outside the filled lines, the report reads as its template."""
    template = docx.Document(str(documents.TEMPLATES / documents.DOKLAD_TEMI))
    before = [p.text for p in template.paragraphs]
    after = [p.text for p in document.paragraphs]
    assert len(before) == len(after)
    for old, new in zip(before, after):
        if "{{" not in old:
            assert new == old


@pytest.mark.parametrize("code, dean, address", [
    ("BSNDIU01", "ДОЦ. Д-Р ВЕНЦИСЛАВ ДЖАМБАЗОВ", "ДЕКАН НА ФДЕНО"),
    ("BSNBUB04", "ДОЦ. Д-Р ИЛИЯ КОЖУХАРОВ", "ДЕКАН НА БАКАЛАВЪРСКИ ФАКУЛТЕТ"),
    ("MSNMBA01", "ПРОФ. Д-Р БОРИС СЕРГИНОВ", "ДЕКАН НА МАГИСТЪРСКИ ФАКУЛТЕТ"),
])
def test_with_the_starter_lists_each_dean_is_addressed_as_the_samples_do(
        tmp_path, code, dean, address):
    # The owner, 12 October 2026: „ДЕКАН НА ФДЕНО“, as the ФДЕНО samples
    # write it; БФ and МФ in full.
    lists.write_starter(tmp_path)
    starter = lists.load(tmp_path)
    rows = [{"position": 1, "application_id": 1, "fn": "F1",
             "name": "Ана Бойкова Вълчева", "topic": "Тема",
             "supervisor": "доц. д-р Петър Иванов"}]
    body = [p.text for p in read(documents.doklad_temi(
        code, rows, "2026-10-12T09:00:00+00:00", starter)).paragraphs]
    assert body[1:3] == [dean, address]


# --- The templates themselves ------------------------------------------------

def test_the_report_template_holds_the_header_and_one_templated_row():
    template = docx.Document(str(documents.TEMPLATES / documents.DOKLAD_TEMI))
    (table,) = template.tables
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    assert rows[0] == REPORT_HEADER
    control = [r for r in rows[1:] if r[0].startswith("{%tr")]
    data_rows = [r for r in rows[1:] if not r[0].startswith("{%tr")]
    assert [r[0] for r in control] == ["{%tr for s in students %}",
                                       "{%tr endfor %}"]
    assert all(not any(r[1:]) for r in control)
    assert data_rows == [["{{ s.position }}", "{{ s.fn }}", "{{ s.name }}",
                          "{{ s.topic }}", "{{ s.supervisor }}", ""]]


@pytest.mark.parametrize("name", TEMPLATES)
def test_no_template_prints_a_year_or_logic(name):
    with zipfile.ZipFile(str(documents.TEMPLATES / name)) as package:
        for part in package.namelist():
            if not part.endswith(".xml"):
                continue
            xml = package.read(part).decode("utf-8")
            for text in re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", xml):
                assert not re.search(r"20\d\d", text), (part, text)
                # Only {{ }} tags, and the report's two {%tr %} rows.
                for tag in re.findall(r"{%.*?%}", text):
                    assert tag in ("{%tr for s in students %}",
                                   "{%tr endfor %}")
    core = docx.Document(str(documents.TEMPLATES / name)).core_properties
    assert core.author == "" and core.last_modified_by == ""
    assert core.comments == ""


# --- The student's downloads -------------------------------------------------

def _student(app, address):
    client = Client(app.test_client())
    with client.raw.session_transaction() as session:
        session["student"] = address.casefold()
        session["student_email"] = address
    return client


def _app(tmp_path):
    data_dir = tmp_path / "data"
    prepare_data_dir(data_dir)
    settings = make_settings(data_dir)
    app = create_app(settings)
    app.config["TESTING"] = True
    return settings, app


def test_a_student_downloads_only_their_own_submitted_forms(tmp_path):
    settings, app = _app(tmp_path)
    insert_application(settings, "thesis@example.org", grad_type="thesis",
                       name1="Ина", name2="Йорданова", name3="Кирова",
                       fn="F200001", programme="BSNADM01", professor="ivanov",
                       topic="Тема на тезата", rationale="ред 1\nред 2")
    insert_application(settings, "exam@example.org", grad_type="exam",
                       name1="Лора", name2="Методиева", name3="Ненова",
                       fn="F200002")
    insert_application(settings, "draft@example.org", submitted=False,
                       grad_type="thesis", name1="Олга", fn="F200003")

    thesis = _student(app, "thesis@example.org")
    reply = thesis.get("/application/forms/zaiavlenie.docx")
    assert reply.status_code == 200
    assert reply.mimetype == documents.MIMETYPE
    assert "attachment" in reply.headers["Content-Disposition"]
    assert reply.headers["Content-Disposition"].isascii()
    text = "\n".join(p.text for p in every_paragraph(read(reply.data)))
    assert "Ина Йорданова Кирова" in text and "Лора" not in text
    reply = thesis.get("/application/forms/prilozhenie2.docx")
    assert reply.status_code == 200
    assert "ред 1\nред 2" in [p.text for p in read(reply.data).paragraphs]

    exam = _student(app, "exam@example.org")
    reply = exam.get("/application/forms/zaiavlenie.docx")
    assert reply.status_code == 200
    text = "\n".join(p.text for p in every_paragraph(read(reply.data)))
    assert "Лора Методиева Ненова" in text and "Ина" not in text
    assert exam.get("/application/forms/prilozhenie2.docx").status_code == 404

    draft = _student(app, "draft@example.org")
    for form in ("zaiavlenie", "prilozhenie2"):
        url = "/application/forms/" + form + ".docx"
        assert draft.get(url).status_code == 404
        assert _student(app, "nobody@example.org").get(url).status_code == 404
        signed_out = Client(app.test_client()).get(url)
        assert signed_out.status_code == 302
        assert signed_out.headers["Location"].endswith("/")
