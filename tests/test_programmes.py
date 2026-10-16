"""The programme list of 13 October 2026 (constitution 2.6 and 5.3): one row
per programme and form, one code each, chosen by students by name only and
shown to staff with its code; and the one-time move of the pilot's records
from the first starting list's codes, as corrections under the owner's
account, once between the service and the reminder run. Everyone is
invented.
"""

import csv
import io
import multiprocessing
import re
import traceback

import docx
import pytest

from superassistant import applications, db, documents, lists, reports
from superassistant.web import StartupError, create_app, prepare
from tests.conftest import (Client, connect, insert_application,
                            insert_check, insert_request, make_account,
                            make_settings, prepare_data_dir)
from tests.test_application_flow import (as_staff, as_student, record_for,
                                         submit_thesis)
from tests.test_documents import every_paragraph

# The owner's rows, as Fable derived them from the owner's instruction of 13
# October 2026 (programmes-proposed-v3.csv).
OWNERS_CODES = [
    "BSNADM01", "BSNADM04", "BSNBUB04", "BSNDIU01", "BSNTUR01", "BSNDIT01",
    "BSNDIT04", "MSNADM21", "MSNALH03", "MSNHRB05", "MSNHRE08", "MSNMBA01",
    "MSNMBA12", "MSNMBA81", "MSNMBA80", "MSNSMQ11", "MSNSMQ09",
]
OLD_LIST = [
    ["АУ-РО", "БП", "Администрация и управление", "РО", "БФ"],
    ["УБП-РО", "БП", "Управление на бизнеса и предприемачество", "РО", "БФ"],
    ["УБП-ДО", "БП", "Управление на бизнеса и предприемачество", "ДО",
     "ФДЕНО"],
    ["АМ-РО", "МП", "Артмениджмънт. Управление на културни институции и "
     "творчески индустрии", "РО", "МФ"],
    ["МОС-РО", "МП", "Мениджмънт за организационно съвършенство", "РО", "МФ"],
    ["МОС-ДО", "МП", "Мениджмънт за организационно съвършенство", "ДО",
     "ФДЕНО"],
    ["УРЧР-РО", "МП", "Управление и развитие на човешките ресурси", "РО",
     "МФ"],
    ["БА-РО", "МП", "Бизнес администрация", "РО", "МФ"],
]
OWNER = "boss"


def write_programmes(lists_dir, rows):
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(lists.HEADERS[lists.PROGRAMMES])
    for row in rows:
        writer.writerow(row)
    (lists_dir / lists.PROGRAMMES).write_text(buffer.getvalue(),
                                              encoding="utf-8")


# --- The list ----------------------------------------------------------------

def test_the_starting_list_is_the_owners_one_row_per_programme_and_form(
        tmp_path):
    lists.write_starter(tmp_path)
    loaded = lists.load(tmp_path)
    assert list(loaded.programmes) == OWNERS_CODES
    pairs = [(p["name"], p["form"]) for p in loaded.programmes.values()]
    assert len(set(pairs)) == len(pairs)          # each programme and form once
    names = " ".join(p["name"] for p in loaded.programmes.values())
    assert "Курсист" not in names and "здравеопазването" not in names
    for code, programme in loaded.programmes.items():
        # The degree from the code's prefix; the faculty by the owner's
        # mapping: ДО to ФДЕНО, РО bachelor's to БФ, РО master's to МФ.
        assert programme["degree"] == {"BSN": "БП", "MSN": "МП"}[code[:3]]
        expected = ("ФДЕНО" if programme["form"] == "ДО" else
                    "БФ" if programme["degree"] == "БП" else "МФ")
        assert programme["faculty"] == expected
    # The header and the loader's rules are those of before.
    header = (tmp_path / lists.PROGRAMMES).read_text(
        encoding="utf-8").splitlines()[0]
    assert header == "код,степен,програма,форма,факултет"


def test_every_old_starter_code_has_a_listed_programme_of_the_same_form(
        tmp_path):
    lists.write_starter(tmp_path)
    loaded = lists.load(tmp_path)
    assert sorted(lists.RENAMED_PROGRAMMES) == sorted(r[0] for r in OLD_LIST)
    for code, degree, name, form, faculty in OLD_LIST:
        new = loaded.programmes[lists.RENAMED_PROGRAMMES[code]]
        assert (new["degree"], new["form"], new["faculty"]) == (
            degree, form, faculty)
        assert new["name"].startswith(name)


# --- Students choose by name; staff see the code ------------------------------

def test_students_choose_by_name_and_never_see_a_code(client, settings,
                                                      outbox):
    as_student(client, "label@example.org")
    page = client.get("/application").get_data(as_text=True)
    options = re.findall(r'<option value="([A-Z0-9]+)"\s*>([^<]*)</option>',
                         page)
    loaded = lists.load(settings.lists_dir)
    assert [code for code, _ in options] == OWNERS_CODES
    for code, text in options:
        programme = loaded.programmes[code]
        assert text == "{} „{}“, {}".format(programme["degree"],
                                            programme["name"],
                                            programme["form"])
    visible = re.sub(r'value="[^"]*"', "", page)
    assert not any(code in visible for code in OWNERS_CODES)
    submit_thesis(client, {"programme": "MSNHRB05"})
    page = client.get("/application").get_data(as_text=True)
    assert "Управление и развитие на човешките ресурси (на български и " \
        "английски език)" in page
    assert not any(code in page for code in OWNERS_CODES)


def test_staff_see_the_applications_code_in_the_table_and_on_its_page(
        client, new_client, settings, outbox):
    as_student(client, "code@example.org")
    submit_thesis(client, {"programme": "MSNSMQ09"})
    record = record_for(settings, "code@example.org")
    viewer = as_staff(new_client(), make_account(settings, "staff1"))
    table = viewer.get("/staff").get_data(as_text=True)
    assert "<th>Програма</th><th>Код</th>" in table
    row = table.split('data-id="{}"'.format(record["id"]))[1].split(
        "</tr>")[0]
    assert re.search(r"<td>МП Мениджмънт за организационно съвършенство, "
                     r"ДО</td>\s*<td class=\"code\">MSNSMQ09</td>", row)
    page = viewer.get("/staff/application/{}".format(record["id"])
                      ).get_data(as_text=True)
    assert '<dt>Код на програмата</dt><dd id="programme-code">MSNSMQ09</dd>' \
        in page
    # The correction form lists every code with its programme.
    hint = page.split("За програма въведете кода ѝ:")[1].split("</p>")[0]
    assert all(code in hint for code in OWNERS_CODES)


# --- The old codes move once, as the owner's corrections ----------------------

def pilot_records(settings):
    """One application per old code, a draft among them, and a kept report."""
    connection = connect(settings)
    try:
        owner = make_account(settings, OWNER, approved=False)
        numbers = {}
        for index, (code, *_rest) in enumerate(OLD_LIST):
            numbers[code] = insert_application(
                settings, "old{}@example.org".format(index),
                submitted=index != 3, grad_type="thesis", programme=code,
                name1="Име", name2="Бащино", name3="Фамилия" + str(index),
                fn="F{}".format(100 + index), topic="Тема",
                professor="ivanov")
        exam = insert_application(settings, "exam@example.org",
                                  grad_type="exam", name1="Изпит")
        with db.transaction(connection):
            connection.execute(
                "INSERT INTO reports (programme, faculty, generated_at, "
                "generated_by, filename) VALUES ('АУ-РО', 'БФ', ?, ?, "
                "'doklad-temi-old.docx')", (db.now(), owner["id"]))
            connection.execute(
                "INSERT INTO report_students VALUES (1, 1, ?, 'F100', 'Име', "
                "'Тема', 'доц. д-р Петър Иванов')", (numbers["АУ-РО"],))
        return owner, numbers, exam
    finally:
        connection.close()


@pytest.fixture
def pilot(tmp_path):
    """A data directory as the pilot's: the old codes in the records, the
    owner's new list file in place, the owner's account named."""
    data_dir = tmp_path / "data"
    prepare_data_dir(data_dir)
    settings = make_settings(data_dir, owner_username=OWNER)
    connection = connect(settings)
    with db.transaction(connection):
        connection.execute("INSERT INTO installation VALUES ('lists', ?)",
                           (db.now(),))
    connection.close()
    owner, numbers, exam = pilot_records(settings)
    return settings, owner, numbers, exam


def test_the_old_codes_move_once_as_corrections_under_the_owners_account(
        pilot):
    settings, owner, numbers, exam = pilot
    app = create_app(settings)
    connection = connect(settings)
    for old, number in numbers.items():
        record = applications.get(connection, number)
        assert record["programme"] == lists.RENAMED_PROGRAMMES[old]
        (correction,) = applications.corrections_of(connection, number)
        assert (correction["field"], correction["old_value"],
                correction["new_value"], correction["username"]) == (
            "programme", old, lists.RENAMED_PROGRAMMES[old], OWNER)
    assert applications.get(connection, exam)["programme"] == ""
    assert applications.corrections_of(connection, exam) == []
    assert [r["programme"] for r in reports.kept(connection)] == ["BSNADM01"]
    assert [r["key"] for r in connection.execute(
        "SELECT key FROM installation ORDER BY key")] == [
        "lists", "programme codes", "sessions list"]
    count = connection.execute("SELECT COUNT(*) FROM corrections").fetchone()[0]
    connection.close()
    assert count == len(numbers)
    # Staff read the move among the application's corrections, and the kept
    # report under the programme's name from the new list.
    viewer = as_staff(Client(app.test_client()), owner)
    page = viewer.get("/staff/application/{}".format(numbers["УРЧР-РО"])
                      ).get_data(as_text=True)
    assert "Програма и форма на обучение: от „УРЧР-РО“ на „MSNHRB05“ ({}, " \
        .format(OWNER) in page
    page = viewer.get("/staff/reports").get_data(as_text=True)
    assert "БП „Администрация и управление“, РО</td>" in page
    # Started again: nothing more moves.
    create_app(settings)
    connection = connect(settings)
    assert connection.execute("SELECT COUNT(*) FROM corrections").fetchone()[
        0] == len(numbers)
    connection.close()


def _prepare_when_released(barrier, data_dir, results):
    try:
        settings = make_settings(data_dir, owner_username=OWNER)
        barrier.wait(timeout=60)
        prepare(settings)
        results.put("ok")
    except BaseException:
        results.put(traceback.format_exc())


@pytest.mark.parametrize("round_number", range(3))
def test_the_service_and_the_reminder_run_move_the_codes_once_between_them(
        pilot, round_number):
    settings, owner, numbers, exam = pilot
    context = multiprocessing.get_context("spawn")
    barrier = context.Barrier(2)
    results = context.Queue()
    workers = [context.Process(target=_prepare_when_released,
                               args=(barrier, str(settings.data_dir), results))
               for _ in range(2)]
    for worker in workers:
        worker.start()
    outcomes = [results.get(timeout=120) for _ in workers]
    for worker in workers:
        worker.join(timeout=60)
    assert outcomes == ["ok", "ok"], "\n".join(outcomes)
    connection = connect(settings)
    try:
        assert connection.execute(
            "SELECT COUNT(*) FROM corrections").fetchone()[0] == len(numbers)
        assert connection.execute(
            "SELECT COUNT(*) FROM installation WHERE key = 'programme codes'"
        ).fetchone()[0] == 1
    finally:
        connection.close()


def test_without_the_owners_account_nothing_moves_and_the_start_stops(pilot):
    settings, owner, numbers, exam = pilot
    unnamed = make_settings(settings.data_dir)
    with pytest.raises(StartupError) as raised:
        create_app(unnamed)
    assert lists.PROGRAMMES in str(raised.value)
    assert "'АМ-РО'" in str(raised.value)             # the first, sorted
    connection = connect(settings)
    assert connection.execute("SELECT COUNT(*) FROM corrections").fetchone()[
        0] == 0
    assert connection.execute("SELECT COUNT(*) FROM installation WHERE key = "
                              "'programme codes'").fetchone()[0] == 0
    connection.close()


def test_a_list_lacking_a_moved_code_stops_the_start(pilot):
    settings, owner, numbers, exam = pilot
    # The update on the old list file: the codes move, the start stops.
    write_programmes(settings.lists_dir, OLD_LIST)
    with pytest.raises(StartupError) as raised:
        create_app(settings)
    assert lists.PROGRAMMES in str(raised.value)
    assert "'BSNADM01'" in str(raised.value)          # the first, sorted
    # The owner's list without one of them: the start names that one.
    lists.write_starter(settings.lists_dir / "fresh")
    rows = [r for r in csv.reader(io.StringIO((
        settings.lists_dir / "fresh" / lists.PROGRAMMES).read_text(
            encoding="utf-8")))][1:]
    write_programmes(settings.lists_dir,
                     [r for r in rows if r[0] != "MSNSMQ09"])
    with pytest.raises(StartupError) as raised:
        create_app(settings)
    assert "'MSNSMQ09'" in str(raised.value)
    assert str(numbers["МОС-ДО"]) in str(raised.value)
    # With the owner's list in place it starts, and nothing moves again.
    write_programmes(settings.lists_dir, rows)
    assert create_app(settings) is not None
    connection = connect(settings)
    assert connection.execute("SELECT COUNT(*) FROM corrections").fetchone()[
        0] == len(numbers)
    connection.close()


# --- The documents print the name from the new list ---------------------------

def test_the_topic_report_and_the_forms_print_the_new_lists_name(
        client, settings, outbox):
    name = ("Управление и развитие на човешките ресурси (на български и "
            "английски език)")
    as_student(client, "hr@example.org")
    submit_thesis(client, {"programme": "MSNHRB05"})
    for form in ("zaiavlenie", "prilozhenie2"):
        reply = client.get("/application/forms/" + form + ".docx")
        text = [p.text for p in every_paragraph(docx.Document(io.BytesIO(
            reply.data)))]
        if form == "zaiavlenie":
            assert "Програма: МП „{}“, РО".format(name) in text
        else:
            assert any("МАГИСТЪРСКА" in line for line in text)
    record = record_for(settings, "hr@example.org")
    staff = make_account(settings, "staff1")
    insert_request(settings, record["id"], answer="confirm")
    insert_check(settings, record["id"], "passed", staff["id"])
    connection = connect(settings)
    loaded = lists.load(settings.lists_dir)
    report_id = reports.generate(connection, settings, loaded, "MSNHRB05",
                                 staff["id"])
    path = reports.path(settings, reports.get(connection, report_id))
    connection.close()
    body = [p.text for p in docx.Document(str(path)).paragraphs]
    assert body[1:3] == ["ПРОФ. Д-Р БОРИС СЕРГИНОВ",
                         "ДЕКАН НА МАГИСТЪРСКИ ФАКУЛТЕТ"]
    assert "Студенти от МП „{}“, изпълнили условията за допускане до " \
        "дипломиране с магистърска теза:".format(name) in body
    assert documents.programme_label("MSNHRB05", loaded) == \
        "МП „{}“, РО".format(name)
